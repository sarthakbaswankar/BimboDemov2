"""Builds 'Set of rules for automation/National_Customers_Rules_Guide.xlsx'.

A plain-language companion to National_Customers_Automation_Rules.xlsx, written
for AR team members and future maintainers. It explains every rule with the
reason behind it, step-by-step logic, real examples from the sample files and
what the person running the process has to check. Rule IDs (G-xx, M-xx, C-xx)
are the same as in the rules workbook and in the 'Rule' column of the outputs.

The code does NOT read this file; parameters are still changed in the
Parameters sheet of National_Customers_Automation_Rules.xlsx.

Run:  python automation/build_rules_guide.py
"""
import sys
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
from config import DEFAULTS, RULES_PATH, to_cell  # noqa: E402

GUIDE_PATH = RULES_PATH.parent / "National_Customers_Rules_Guide.xlsx"

# plain styling: light grey header, common rules lightly shaded
HEADER_FILL = PatternFill("solid", fgColor="F2F2F2")
COMMON_FILL = PatternFill("solid", fgColor="F7F7F7")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# ---------------------------------------------------------------------------------
# 1. Start here
# ---------------------------------------------------------------------------------
START = [
    ("What is this file?",
     "A plain-language explanation of every rule used to apply Metro and Costco payments. It tells you what each rule "
     "does, why we do it, how it works step by step, a real example, and what you have to check."),
    ("Who is it for?",
     "Anyone who runs, reviews or maintains the Metro / Costco cash application: AR team members, new joiners, "
     "reviewers and developers changing the automation."),
    ("How is it different from the rules workbook?",
     "National_Customers_Automation_Rules.xlsx (same folder) is the short technical version and its Parameters sheet "
     "is read by the automation. This guide is the long, explained version. Rule IDs are the same in both files."),
    ("Rule IDs",
     "G-xx = common rule (Metro and Costco). M-xx = Metro only. C-xx = Costco only. "
     "The same IDs appear in the 'Rule' column of the output workbooks, so you can look up why a line got its status."),
    ("Colours",
     "Light grey rows = rule that is common to Metro and Costco (see the 'Common rule' column). "
     "White rows = rule that applies to that customer only."),
    ("Suggested reading order",
     "1 Glossary  >  2 Process_Overview  >  3 Decision_Flow  >  4 Common_Rules  >  5 Metro_Rules / Costco_Rules  >  "
     "6 Worked_Examples  >  7 Output_Guide  >  8 Daily_Checklist"),
    ("Changing a rule",
     "Values such as tolerances, bill-to accounts or transaction type names: edit the Parameters sheet of "
     "National_Customers_Automation_Rules.xlsx (see Parameters_Explained here). Changing the logic itself needs a code "
     "change in the automation folder; update both rule files afterwards."),
    ("Sample files used for examples",
     "Metro: CAN_RMT_METRODE-CA6-RMT_2252_202608271334.csv (control 2000238539, 27-Aug-2026, $164,986.02). "
     "Costco: Costco Payment.PDF (payment 1509518, 01-Sep-2026, $898,467.62)."),
    ("Generated", date.today().isoformat()),
]

# ---------------------------------------------------------------------------------
# 2. Glossary
# ---------------------------------------------------------------------------------
GLOSSARY = [
    ("Remittance / payment advice", "The customer's list of what a payment covers: invoice numbers and amounts. "
     "Metro sends a CSV, Costco a PDF.", "Metro CSV 164,986.02; Costco PDF 898,467.62"),
    ("Open Item Report", "Oracle report of everything still open (not fully paid) on the customer accounts: invoices, "
     "credit memos, deductions. Run every morning before processing.", "Open Item Reports folder"),
    ("Open amount / Due Remaining", "How much of an Oracle document is still unpaid.", ""),
    ("Bill-to", "The customer account an Oracle document is billed to. For Costco it decides the region.",
     "60463822 = Costco Western"),
    ("Ship-to", "The store/location an invoice was delivered to. Used to find the right invoice when numbers differ.",
     "60529764 = Costco store 656"),
    ("Control number (Metro)", "Metro's payment number, shown in the REMITTANCE SUMMARY row.", "2000238539"),
    ("Payment # (Costco)", "Costco's payment number at the top of the PDF.", "1509518"),
    ("Gross / Discount / Net (Metro)", "Gross = invoice amount; Discount = 0.25% Metro takes on invoices; "
     "Net = what Metro actually paid (Gross - Discount).", "125.81 - 0.31 = 125.50"),
    ("Match key / R12", "The cleaned-up invoice number used for matching (see G-01). R12 = 'right 12 digits', the "
     "Costco version.", "1261732219883 -> 261732219883"),
    ("Matched", "The payment line was found in Oracle and the amount agrees.", ""),
    ("Resolved (Costco)", "The amounts did not agree at first, but the engine found the other Oracle document(s) "
     "that explain the difference, so the difference is now 0.", "656 + credit memo"),
    ("Variance (K)", "Paid by the customer minus open in Oracle, for one document.", "8,980.62 - 8,966.80 = 13.82"),
    ("Net variance (L)", "The variance of the whole payment line after all matching, with the sign flipped. "
     "Positive = customer paid less (short pay); negative = customer paid more (overpay).", "L = -13.82 -> overpay"),
    ("Deduction", "Money the customer took off the payment for its own reasons (promotions, claims, returns). "
     "We create it in Oracle so the trade team can work it.", "Metro 0408274209 -160,000.00"),
    ("RV (Costco)", "A type of Costco deduction whose reference starts with RV.", "RV000551013919 -8.75"),
    ("NF - Not Found", "A number the customer paid or deducted that does not exist in Oracle. Booked as NF-<number>.",
     "NF-26506520282"),
    ("SB INV - Should-be invoice (Metro)", "The invoice exists in Oracle but is already closed (paid or credited), so "
     "it is not on the open items. Booked as SB INV-<number>.", "SB INV-26921620976-1"),
    ("SP - Short pay (Costco)", "Customer paid less than we can explain, $11 or more. Booked as "
     "<store>-<MMDDYY>-SP (a debit, the customer still owes it).", "656-080326-SP 100.28"),
    ("OP - Overpay (Costco)", "Customer paid more than we can explain, $11 or more. Booked as <store>-<MMDDYY>-OP "
     "(a credit on the account).", "542-081426-OP -37.92"),
    ("WO - Write-off", "A small difference we absorb instead of creating a transaction. Metro up to $10, "
     "Costco up to $11 (either sign).", "226603721412 8.16"),
    ("DISC (Metro)", "Debit memo for the 0.25% discount Metro took, DISC-<last 7 of control>.", "DISC-0238539 2,021.69"),
    ("NAC", "The balancing line that makes every upload total 0 and that the receipt applies against. "
     "Metro: NAC<last 7 of control>. Costco: NAC<payment>-<region letter>.", "NAC0238539; NAC1509518-Q"),
    ("Upload / AutoInvoice / FBDI", "The file used to create new transactions (deductions, NF, SB, SP, OP, DISC, NAC) "
     "in Oracle. The automation prepares the lines on the Upload sheet; the FBDI file itself is not generated yet.", ""),
    ("Receipt", "The cash receipt in Oracle for the money received. It is applied to the matched invoices and the NAC "
     "line(s). Its amount = remittance total.", ""),
    ("DR / CR transaction types", "CACBC-DEDUC-MGMT-DR for positive amounts (debit, customer owes us). "
     "CACBC-DEDUC-MGMT-CR for negative amounts (credit, we owe / hold for the customer).", ""),
    ("Review sheet", "Lines the automation was not fully sure about. A person must check them before upload.", ""),
]

# ---------------------------------------------------------------------------------
# 3. Process overview
# ---------------------------------------------------------------------------------
PROCESS = [
    ("Both", "0", "Run the Open Item Report", "Run it the same morning before processing (status Open). Metro: the "
     "collector report. Costco: all regions in one file.", "Manual", "-", "G-11"),
    ("Metro", "1", "Read the remittance CSV", "Reads the control number, total and every invoice line. Region comes "
     "from the file name (CA6 = Ontario, CA5 = Quebec). Checks that the lines add up to the total.", "Automation",
     "Raw_Remittance, Control_Checks", "M-01, M-02"),
    ("Metro", "2", "Match each line to Oracle", "Looks up each invoice (leading zeros ignored) on the Metro, Metro "
     "Richelieu and Adonis bill-tos. Same amount = Matched.", "Automation", "Formatted", "M-03, G-01"),
    ("Metro", "3", "Classify what is not found", "Short numbers = deductions; small ones = write-off; "
     "invoice-sized numbers = NF or SB.", "Automation", "Deductions, NF_SB, WriteOffs", "M-05 to M-11"),
    ("Metro", "4", "Build the upload", "Deductions, NF, SB (sign flipped), the DISC memo and the NAC balancing line.",
     "Automation", "Upload", "M-12 to M-14"),
    ("Metro", "5", "Build the receipt", "Matched invoices + NAC = remittance total.", "Automation", "Receipt",
     "M-15, M-16"),
    ("Metro", "6", "Check and confirm", "Control_Checks all PASS; confirm the provisional NF lines on Review in Oracle "
     "(NF or SB?).", "Manual", "Control_Checks, Review", "M-09"),
    ("Costco", "1", "Read the PDF", "Reads every line on every page; the lines must add up exactly to Remittance "
     "TOTAL or the run stops.", "Automation", "Raw_Remittance, Control_Checks", "C-01, C-02"),
    ("Costco", "2", "Combine duplicates", "Same invoice number twice (invoice + Costco's credit) becomes one line.",
     "Automation", "Formatted", "C-04"),
    ("Costco", "3", "Match to Oracle", "Looks up each invoice by its R12 number. Same amount = Matched.",
     "Automation", "Atlantic / Western / Ontario / Quebec", "C-03"),
    ("Costco", "4", "Resolve differences", "Tries credit memos, neighbour invoice numbers and combinations of "
     "invoices for the same store and day until the difference is 0.", "Automation", "Resolution_Log",
     "C-11 to C-14"),
    ("Costco", "5", "Classify what is left", "Small difference = WO; bigger = SP/OP; not in Oracle = NF; deductions "
     "go to Ontario.", "Automation", "Regional sheets, WriteOffs", "C-05, C-06, C-09, C-15, C-16"),
    ("Costco", "6", "Build upload and receipt", "Per region: SP, OP, NF, deductions + NAC line. Receipt = matched "
     "documents + NAC lines.", "Automation", "Upload, Receipt", "C-17 to C-19"),
    ("Costco", "7", "Check and confirm", "Control_Checks all PASS; work the Review sheet (large SP/OP, engine matches, "
     "partial matches).", "Manual", "Control_Checks, Review", "C-14, C-16"),
    ("Both", "Later", "Load into Oracle", "Copy Upload to the AutoInvoice/FBDI template, run Load Interface File and "
     "Import AutoInvoice, create and post the receipt, create write-offs.", "Manual (out of scope now)", "-",
     "G-12, G-13"),
]

# ---------------------------------------------------------------------------------
# 4. Decision flow
# ---------------------------------------------------------------------------------
FLOW = [
    ("Metro", "1", "Is the invoice number open in Oracle on a Metro-group bill-to (leading zeros ignored)?",
     "Go to 2", "Go to 4", "M-03"),
    ("Metro", "2", "Does the gross amount equal the Oracle open amount?", "MATCHED: goes on the receipt",
     "Go to 3", "M-03"),
    ("Metro", "3", "Is the difference between -10.00 and +10.00?",
     "WO: invoice applied at the Oracle amount, difference written off",
     "REVIEW: applied at the paid amount, investigate", "M-08"),
    ("Metro", "4", "Does the number have fewer than 10 digits (after removing leading zeros)?",
     "It is a DEDUCTION: go to 5", "It is an invoice number: go to 6", "M-05"),
    ("Metro", "5", "Is the net amount between -10.00 and +10.00?", "WO: not uploaded",
     "DEDUCTION: uploaded at net x -1", "M-06, M-07"),
    ("Metro", "6", "Does the invoice exist in Oracle but closed? (closed-transactions export, or SB INV-<number> "
     "already in Oracle)", "SB: SB INV-<number> (+ -1/-2 if already used)",
     "NF: NF-<number> (provisional and on Review if there is no export)", "M-09, M-10"),
    ("Costco", "1", "Does the same invoice number appear more than once?", "Add the lines into one line",
     "Keep as is", "C-04"),
    ("Costco", "2", "Is the invoice (R12) open in Oracle?", "Go to 3", "Go to 7", "C-03"),
    ("Costco", "3", "Does the paid amount equal the open amount?", "MATCHED", "Go to 4", "C-03"),
    ("Costco", "4", "Can the engine explain the difference? Tried in order: credit memo (same store, day, amount) > "
     "invoice number +1 / +2 / -1 > 1-4 invoices of the same store and day (then the day before) > "
     "partial match", "RESOLVED (if the difference reaches 0)", "Go to 5 with what is left", "C-11 to C-14"),
    ("Costco", "5", "Is what is left (L) between -11.00 and +11.00?", "WO: not uploaded", "Go to 6", "C-15"),
    ("Costco", "6", "Is L positive?", "SP <store>-<MMDDYY>-SP (debit)", "OP <store>-<MMDDYY>-OP (credit)", "C-16"),
    ("Costco", "7", "Does it look like a deduction? (starts with RV, 65, 14 or 604, or has 13+ digits)",
     "DEDUCTION in Ontario: RV under its own number, others as NF-<number>", "Go to 8", "C-05, C-06"),
    ("Costco", "8", "Can the engine find Oracle documents for that store and date?",
     "Treated like step 4-6 and flagged on Review", "NF: NF-<number>", "C-09"),
]

# ---------------------------------------------------------------------------------
# 5. Rules (ID, common link, rule, what, why, how, example, automation, you check, where)
# ---------------------------------------------------------------------------------
COMMON_RULES = [
    ("G-01", "Both", "Match on a cleaned-up invoice number",
     "Before comparing a customer's invoice number with Oracle, both numbers are cleaned so that the same invoice "
     "always looks the same.",
     "Customers and Excel add or drop leading zeros, and Costco sometimes adds an extra digit in front. Without "
     "cleaning, real matches would show as N/A (not found).",
     "Metro: remove all leading zeros on both sides.\n"
     "Costco: take the last 12 digits and put zeros in front until it is 12 digits long (the old R12 column).",
     "Metro: CSV 026593522226 and Oracle 26593522226 both become 26593522226, so they match.\n"
     "Costco: PDF 1261732219883 becomes 261732219883; 26866221906 becomes 026866221906.",
     "Done automatically. There is no R12 column to maintain any more.",
     "Nothing, normally. If a line you expected to match shows NF, check how the customer typed the number.",
     "Metro: Formatted > Match Key. Costco: regional sheets > R12."),
    ("G-02", "Both", "One Oracle document can only be used once",
     "Each open invoice or credit in Oracle can explain only one payment line.",
     "Using the same invoice twice would apply the same money twice and the receipt would not balance in Oracle.",
     "1. Exact matches are reserved first.\n2. Every document attached later is marked as used.\n"
     "3. Later searches skip used documents.",
     "Costco 226307921596 is attached to payment line 226307921597 (store 656) and cannot be used by any other line.",
     "Automatic. Control 'No open item used twice' / 'No open item applied twice'.",
     "If that control is FAIL, do not upload. Report it.",
     "Control_Checks; Costco Resolution_Log."),
    ("G-03", "Both", "The upload amount is the opposite sign of the remittance amount",
     "What we create in Oracle has the opposite sign of what the customer deducted or paid on the remittance.",
     "If the customer deducts $100 (-100 on the remittance), the customer still owes it, so we create +100 (debit). "
     "If the customer pays $50 for an invoice we cannot find (+50), we hold -50 (credit) on the account.",
     "Upload amount = remittance amount x -1.\n"
     "Costco SP/OP use the net variance L, which is already flipped.",
     "Metro deduction 0408274209: -160,000.00 becomes +160,000.00.\n"
     "Metro NF-26543919847: paid +1,430.03 becomes -1,430.03.\n"
     "Costco NF-6518087993: -43,591.80 becomes +43,591.80.",
     "Automatic.", "Spot-check a few lines on the Upload sheet.",
     "Upload > Transaction Line Amount; Metro Formatted > Upload Amount; Costco regional sheets > Upload Amount."),
    ("G-04", "Both", "Transaction type comes from the sign",
     "Every upload line gets a transaction type based on its sign.",
     "Oracle needs to know whether the new transaction is something the customer owes (debit) or something we owe "
     "/ hold for the customer (credit).",
     "Amount above 0: CACBC-DEDUC-MGMT-DR (debit). Metro debits also get payment terms Net 15.\n"
     "Amount below 0: CACBC-DEDUC-MGMT-CR (credit).",
     "Metro deduction +160,000.00: DR, Net 15. Metro NAC0238539 -571,867.58: CR.",
     "Automatic. Type names and terms come from the Parameters sheet.",
     "Note: an early version of the Costco notes said 'positive' for both. Correct is DR = positive, CR = negative.",
     "Upload > Transaction Type Name, Payment Terms."),
    ("G-05", "Both", "Every upload balances to 0 with a NAC line",
     "Each group of new transactions gets one extra line, the NAC, so that the group adds up to exactly 0.",
     "Accounting: every debit needs an equal credit, so the clearing account ends at 0. The NAC is also the "
     "document the receipt applies against, which makes the receipt equal the money received.",
     "1. Add up all upload lines of the group.\n2. NAC = that total with the opposite sign.\n"
     "Metro: one NAC per receipt, named NAC + last 7 digits of the control number.\n"
     "Costco: one NAC per region, named NAC<payment #>-<A/W/O/Q>.",
     "Metro: deductions +572,219.34, NF/SB -2,373.45, DISC +2,021.69 = 571,867.58, so NAC0238539 = -571,867.58.\n"
     "Costco Quebec: one OP of -37.92, so NAC1509518-Q = +37.92.",
     "Automatic. Control 'Upload sums to 0' (one per NAC for Costco).",
     "Check the control is PASS.", "Upload (Category = NAC); Receipt."),
    ("G-06", "Both", "Upload dates are written YYYY-MM-DD",
     "Transaction and accounting dates on the upload use the format year-month-day.",
     "Any other format does not upload properly in Oracle (Metro walkthrough).",
     "Metro: the receipt date = the day the money / file was received (default: the date in the file name).\n"
     "Costco: the working date (default: the date on the PDF). Both can be overridden when running.",
     "Metro file ..._202608271334.csv: 2026-08-27. Costco PDF 'Date : 09/01/2026': 2026-09-01.",
     "Automatic.", "If you process a file on a later day, run with --receipt-date / --work-date.",
     "Upload > Transaction Date, Accounting Date."),
    ("G-07", "Both", "What goes into each upload line",
     "Each upload line is filled the way the AutoInvoice template expects it.",
     "These are the fields the team fills by hand today; the rest of the template stays fixed.",
     "1. The reference is written 3 times: Transaction Number, Transaction Line Description, Line Transactions "
     "Flexfield Segment 1.\n2. The bill-to is written twice: account number and site number.\n"
     "3. Fixed values from Parameters: business unit (CACBC), batch source, line type (LINE), currency (CAD).",
     "Metro line 408274209: number, description and flexfield are all 408274209; bill-to 60477295 in both columns.",
     "Automatic.", "Confirm the fixed values match the team's template (Open question Q10).",
     "Upload sheet (columns after 'Line Transactions Flexfield Segment 1' are helper columns, not for the FBDI)."),
    ("G-08", "Both", "References must be unique: add -1, -2",
     "A reference (NF-..., SB INV-..., SP/OP, NAC...) that already exists gets a suffix.",
     "Oracle rejects a transaction number that is already used (AutoInvoice reject).",
     "1. Before using a name, check the Open Item Report and the names created earlier in the same run.\n"
     "2. If taken, try <name>-1, then <name>-2, and so on.",
     "SB INV-26921620976 is already open in Oracle, so the new one becomes SB INV-26921620976-1.",
     "Automatic, but it can only see open items. A closed transaction with the same name is not in the report.",
     "If AutoInvoice still rejects a line as a duplicate, change the suffix (e.g. add today's date) and re-import.",
     "Reference column on all sheets; Note column says when a suffix was added."),
    ("G-09", "Both", "The receipt equals the money received",
     "The receipt amount is the remittance total. Everything applied on the receipt plus the write-offs must add up "
     "to it.",
     "If the receipt does not balance, Oracle shows an Amount Due and the cash is not fully applied.",
     "Receipt lines = every matched Oracle document (at its Oracle amount) + the NAC line(s).\n"
     "Deductions, NF, SB, SP and OP are not on the receipt one by one: they are inside the NAC.",
     "Metro: matched 736,853.60 + NAC -571,867.58 = 164,986.02 = remittance total.",
     "Automatic. Control 'Receipt references + write-offs = remittance total'.",
     "Check the control is PASS before creating the receipt.", "Receipt sheet; Control_Checks."),
    ("G-10", "Both", "The remittance must add up before anything else",
     "The lines read from the file must add up to the total printed by the customer.",
     "If they don't, a line was missed or misread and every later step would be wrong.",
     "Metro: sum of net = REMITTANCE SUMMARY total, and gross - discount = net.\n"
     "Costco: sum of PDF lines = Remittance TOTAL, exactly. If not, the run stops (use --force only to "
     "investigate).",
     "Metro: 167,007.71 - 2,021.69 = 164,986.02. Costco: 251 lines = 898,467.62.",
     "Automatic.", "If Costco stops, look at the Review sheet 'parse exception' rows: the PDF layout may have "
     "changed.", "Control_Checks (first rows)."),
    ("G-11", "Both", "Run the Open Item Report the same morning",
     "The open items must be a fresh extract, taken before you apply this payment.",
     "If it is old, new invoices are missing and show as NF. If it was taken after other cash was applied, invoices "
     "may already be closed and show as NF or variances.",
     "Status = Open. Metro: the collector report (the automation itself keeps only the Metro-group bill-tos).\n"
     "Costco: all regions from / to customer account in one file.",
     "Metro report runs up to 27-Aug for the 27-Aug file; Costco up to 31-Aug for the 1-Sep payment.",
     "Not automated.", "Run it every day before processing.", "-"),
    ("G-12", "Both", "No FBDI file is generated yet",
     "The automation stops at the Upload and Receipt sheets.",
     "Agreed scope for the first phase.",
     "The Upload sheet uses AutoInvoice (RA_INTERFACE_LINES_ALL) column names so it can be copied into the template.",
     "", "Not automated.", "Copy the Upload columns into the AutoInvoice template and run the jobs as today.", "-"),
    ("G-13", "Both", "Small differences are written off",
     "A difference within the customer's tolerance is not uploaded. It is written off on the receipt instead.",
     "Creating and chasing tiny transactions costs more than the amount. Canada allows a tolerance.",
     "Metro: deduction net amount (or the difference on a matched invoice) between -10.00 and +10.00.\n"
     "Costco: what is left (L) after all matching, between -11.00 and +11.00.\n"
     "Both include the limit itself. Write-offs are not uploaded and not in the NAC.\n"
     "Never written off: NF lines, SB lines, Costco RV deductions.",
     "Costco 226603721412: paid 1,642.68 (1,650.84 - 8.16) against 1,650.84 open, L = 8.16, so WO.\n"
     "Costco 226860621945: L = 10.00, so WO. Metro: none in the 27-Aug file.",
     "Automatic. Lists them on WriteOffs and checks they are within tolerance.",
     "After posting the receipt in Oracle, if Amount Due is not 0: Actions > Create Write-off > Bad Debt "
     "(receivable activity) for the write-off amount, then add the final invoice.",
     "WriteOffs sheet; Receipt check in Control_Checks."),
    ("G-14", "Both", "Not-found items are booked as NF-<number>",
     "A number the customer paid or deducted that cannot be found in Oracle is created as NF-<number>.",
     "We must account for every dollar received; the NF line holds it on the account until someone researches it.",
     "1. Reference = NF- + the number (leading zeros removed by default).\n"
     "2. Amount = remittance amount x -1.\n3. Never written off, whatever the size.",
     "Metro NF-26506520282 (paid 1,693.27, upload -1,693.27). Costco NF-226603721414 (paid 5,109.00, upload "
     "-5,109.00).",
     "Automatic.", "Metro: confirm each provisional NF in Oracle (see M-09).", "NF_SB (Metro); regional sheets "
     "(Costco)."),
    ("G-15", "Both", "Invoice numbers on the receipt are 12 digits",
     "Oracle invoice numbers are 12 digits; the Open Item Report export drops the leading 0 of Ontario invoices.",
     "The receipt would not find the invoice in Oracle with 11 digits.",
     "An 11-digit number gets a 0 in front and is kept as text. Other lengths (deductions like RV..., 65024) are left "
     "as they are.",
     "26593522226 becomes 026593522226.",
     "Automatic. Costco control 'Receipt invoice numbers are 12 digits'.",
     "Do not let Excel turn these columns into numbers when copying (paste as text).",
     "Receipt > Reference Value; Oracle Trx columns."),
    ("G-16", "Both", "Upload lines are sorted from lowest to highest amount",
     "The largest credits are at the top, the largest debits at the bottom.",
     "The team's practice; it also makes the last (largest) invoice easy to find when a write-off is needed.",
     "Sort by Transaction Line Amount, ascending.", "Metro: NAC0238539 -571,867.58 first ... deduction "
     "152960368 +236,337.20 last.", "Automatic.", "-", "Upload sheet."),
]

METRO_RULES = [
    ("M-01", "Metro only", "Read the Metro remittance CSV",
     "Metro sends one CSV per payment. Column 1 says what each row is: REMITTANCE SUMMARY (payment total), "
     "REMITTANCE (invoice lines) and other record types.",
     "Reading the file directly replaces copying it into the template and moving the Net column in front of Gross.",
     "1. Region from the file name: CA6 = Ontario, CA5 = Quebec (Metro sends two separate payments).\n"
     "2. Control number and total from REMITTANCE SUMMARY.\n3. Store, date, invoice, gross, discount, net from "
     "each REMITTANCE row.",
     "CAN_RMT_METRODE-CA6-...csv: Ontario, control 2000238539, 1,428 lines.",
     "Automatic.", "Use the right file (CA5 or CA6). Use --region if the file name has no CA code.",
     "Raw_Remittance; Summary."),
    ("M-02", "G-10", "The Metro file must add up",
     "Sum of the net amounts = payment total, and gross - discount = net.",
     "Proves nothing was lost when reading the file.", "Two checks on the Control_Checks sheet.",
     "167,007.71 - 2,021.69 = 164,986.02.", "Automatic.", "Both controls PASS.", "Control_Checks."),
    ("M-03", "G-01, G-02", "Match each line to an open Metro-group invoice",
     "Each invoice is looked up in the open items of the Metro group only: Metro Ontario (60477295, 60463677), "
     "Metro Richelieu (60503130, 60453574, 60527663, 9201257) and Adonis (60477237, 60503156).",
     "Metro also pays invoices billed to Adonis and a second Metro Ontario account. The collector report also "
     "contains Sobeys, Walmart, etc., which must never be matched to a Metro payment.",
     "1. Clean the number (G-01).\n2. Look it up on the Metro-group bill-tos.\n"
     "3. Compare the GROSS amount with the Oracle open amount.\n4. Equal = Matched (goes on the receipt).",
     "026593522226: gross 125.81, Oracle open 125.81, so Matched (Metro paid net 125.50 after a 0.31 discount).\n"
     "1,380 lines matched, 736,853.60, all with 0 difference.",
     "Automatic. The bill-to list is in Parameters (GROUP_BILLTOS).",
     "31 matched invoices are on Adonis / the second Metro account. Confirm the Ontario receipt can be applied to them "
     "(Open question Q13).", "Formatted > Oracle Trx, Oracle Amount, Oracle Bill-to, Variance."),
    ("M-04", "Metro only", "The 0.25% discount: always use Net for deductions",
     "Metro takes a 0.25% discount on every invoice, never on credits or deductions.",
     "Gross includes the discount. Using gross for a deduction would put the discount in twice and the receipt "
     "would not balance.",
     "Invoices are matched on gross (that is what Oracle has). Deductions use net. The total discount goes into the "
     "DISC memo (M-12).", "125.81 x 0.25% = 0.31. Total discount in the file: 2,021.69.",
     "Automatic.", "-", "Formatted > Gross, Discount, Net."),
    ("M-05", "Metro only", "Short numbers are deductions",
     "A number that is not in Oracle and has fewer than 10 digits (after removing leading zeros) is a deduction "
     "Metro took.",
     "Real Metro invoices are 10-12 digits long. Deduction numbers are shorter and usually start with 1 or 4.",
     "1. Not found in Oracle.\n2. Remove leading zeros.\n3. Fewer than 10 digits = deduction. The reference is the "
     "number itself.",
     "0152960344 -> 152960344 (9 digits): deduction. 0408274209 -> 408274209: deduction. 60122271 (8 digits): "
     "deduction. 41 deductions in the file: 36 negative, 5 positive, total -572,219.34.",
     "Automatic.", "-", "Deductions sheet."),
    ("M-06", "G-03, G-04", "Deductions are uploaded at net x -1",
     "Every deduction outside the write-off range is created in Oracle with the opposite sign.",
     "A negative deduction (Metro kept money) becomes a debit the trade team can work; a positive one becomes a "
     "credit.",
     "Upload amount = net x -1. Positive = DR with Net 15 terms; negative = CR.",
     "0408274209 -160,000.00 -> +160,000.00 DR. 0408274235 +16,148.06 -> -16,148.06 CR.",
     "Automatic.", "-", "Deductions > Upload Amount; Upload sheet."),
    ("M-07", "G-13", "Deductions of $10 or less are written off",
     "A deduction whose net amount is between -10.00 and +10.00 (inclusive) is not uploaded.",
     "Canada's $10 tolerance.", "Status WO. Not on Upload, not in the NAC. Listed on WriteOffs.",
     "None in the 27-Aug file (smallest deduction is 15.95).", "Automatic.",
     "Create the write-off in Oracle after posting the receipt (G-13).", "WriteOffs sheet."),
    ("M-08", "G-13", "Small difference on a matched invoice",
     "If an invoice is found but the paid gross differs from the Oracle amount by $10 or less, the invoice is "
     "applied in full and the difference is written off. A bigger difference goes to Review.",
     "Same $10 tolerance as deductions.",
     "Difference within +/-10: WO (receipt uses the Oracle amount).\n"
     "Above: Review (receipt uses the amount Metro paid, so it still balances).",
     "None in the 27-Aug file (all 1,380 matched with 0 difference).", "Automatic.",
     "Investigate any Review line with a difference before upload.", "Formatted > Variance; Review."),
    ("M-09", "G-14", "Invoice numbers not in Oracle are NF",
     "An invoice-sized number (10 digits or more) that cannot be found in Oracle is booked as NF-<number>.",
     "Metro paid (or credited) something we have no record of. The NF line holds the money on the account.",
     "1. If a closed-transactions export was given (--closed-trx) and the invoice is not in it: NF.\n"
     "2. Without the export the automation cannot tell NF from SB, so it uses NF and puts the line on Review.\n"
     "3. Amount = gross x -1 (DISC_BASIS = TOTAL) or net x -1 (MATCHED).",
     "26543919847 paid 1,430.03 -> NF-26543919847, upload -1,430.03 (credit).\n"
     "26534918561 credit -626.44 -> NF-26534918561, upload +626.44 (debit).",
     "Automatic, but provisional without the export.",
     "Search each Review NF in Oracle. If the invoice exists (closed), change it to SB INV-<number>.",
     "NF_SB sheet; Review."),
    ("M-10", "G-08, G-03", "Invoices that exist but are closed are SB INV",
     "If the invoice exists in Oracle but is already closed (paid or credited), it is booked as SB INV-<number> "
     "('should be invoice').",
     "Metro is paying or crediting something Oracle shows as settled, so there is nothing open to apply to.",
     "1. Found in the closed-transactions export, or an SB INV-<number> already exists in Oracle.\n"
     "2. If SB INV-<number> already exists, add -1, -2 (G-08).\n"
     "3. Positive line: credit memo. Negative line: debit memo.",
     "26921620976 paid 74.29. SB INV-26921620976 (74.29, 13-Aug) is already open in Oracle, so the new one is "
     "SB INV-26921620976-1 with upload -74.29.",
     "Automatic.", "-", "NF_SB sheet."),
    ("M-11", "G-08", "Same invoice twice",
     "Metro sometimes sends an invoice it already paid, by mistake or to take a partial deduction / penalty.",
     "Paying an invoice twice must not be applied twice.",
     "The second occurrence is treated like a closed invoice (SB, with a -1/-2 suffix) and listed on Review.",
     "None in the 27-Aug file.", "Automatic.", "Check with the team whether it is a mistake or a penalty.",
     "Review; Note column."),
    ("M-12", "Metro only", "The discount is booked as DISC-<last 7 of control>",
     "A debit memo for the discount Metro took.",
     "The invoices are applied at gross, but Metro paid net; the DISC memo covers the difference.",
     "DISC_BASIS = TOTAL (default, as in the handwritten notes): the whole discount of the file, and NF/SB use "
     "gross.\nDISC_BASIS = MATCHED: discount on matched invoices only, and NF/SB use net. Both give the same NAC.",
     "DISC-0238539 = +2,021.69.", "Automatic.", "Confirm which basis the team uses (Open question Q2).",
     "Upload (Category DISC); Summary."),
    ("M-13", "G-05", "NAC<last 7 of control> balances the Metro upload",
     "One balancing credit (or debit) per payment, and the receipt number.",
     "Makes the upload 0 and lets the receipt balance to the money received.",
     "NAC amount = -(deductions + NF/SB + DISC).",
     "NAC0238539 = -(572,219.34 - 2,373.45 + 2,021.69) = -571,867.58.", "Automatic.", "-",
     "Upload; Receipt; Summary."),
    ("M-14", "G-06, G-07, G-16", "Metro upload layout",
     "How the Metro upload lines are filled and ordered.", "Matches the team's AutoInvoice practice.",
     "Sorted from largest credit to largest debit. Positive lines get Net 15. Date = receipt date. Ontario and "
     "Quebec files can be uploaded together (one file).",
     "50 lines: 41 deductions, 7 NF/SB, DISC, NAC.", "Automatic.", "-", "Upload."),
    ("M-15", "G-09", "Metro receipt",
     "One receipt per payment (Ontario and Quebec separately).", "Applies the cash in Oracle.",
     "Receipt number = NAC name. Amount = remittance total. Customer account: Ontario 60477295, Quebec 60503130. "
     "Batch 'Metro Ontario YYMMDD'. Method from Parameters. Lines: every matched invoice at its Oracle amount + the "
     "NAC; sequence 1, type Transaction Number. NF/SB lines are NOT on the receipt.",
     "NAC0238539, 164,986.02, batch 'Metro Ontario 260827'.", "Automatic.",
     "Confirm the receipt method name (Q14) and account numbers (Q3).", "Receipt sheet."),
    ("M-16", "G-15", "12-digit invoice numbers on the receipt",
     "Ontario invoices are 12 digits starting with 0; Quebec invoices 12 digits starting with 3.",
     "The export drops the leading 0.", "11 digits get a leading 0.", "26593522226 -> 026593522226.",
     "Automatic.", "Paste as text.", "Receipt > Reference Value."),
    ("M-17", "G-13", "Write-off in Oracle after posting (manual)",
     "After the receipt is uploaded and posted, query it and check Amount Due.",
     "Write-offs were left out of the upload, so the receipt shows them as Amount Due.",
     "Amount Due not 0: Actions > Create Write-off > Bad Debt (receivable activity). Positive remaining: positive "
     "write-off. Negative remaining: negative write-off, then add the final (largest) invoice.",
     "", "Not automated (out of scope now).", "Do it for every line on the WriteOffs sheet.", "WriteOffs."),
    ("M-18", "Metro only", "Other record types",
     "DSD ADJ, MARKET ADJUSTMENTS and COUPON rows can appear in the file.",
     "Their treatment has not been described by the team yet.",
     "Copied to an Other_Records sheet if present.", "None in the sample.", "Captured only.",
     "Tell the team if they appear.", "Other_Records (only when present)."),
]

COSTCO_RULES = [
    ("C-01", "Costco only", "One payment advice PDF per day for all regions",
     "Costco sends one PDF per day covering Atlantic, Western, Ontario and Quebec.",
     "Everything for the day is processed in one run and split into regions afterwards.",
     "The PDF is read directly. Weekdays 6-10 pages; Mondays 22-25 pages (Friday to Sunday).",
     "Payment 1509518: 9 pages, 251 lines.", "Automatic.", "-", "Raw_Remittance; Summary."),
    ("C-02", "G-10", "The PDF must add up exactly",
     "The lines read from the PDF must equal Remittance TOTAL to the cent.",
     "Replaces copy-paste / text-to-columns and the manual 'did anything land in the wrong column?' check.",
     "If the total is different, the run stops.", "251 lines = 898,467.62.", "Automatic.",
     "If it stops, check the Review sheet for 'parse exception' rows.", "Control_Checks."),
    ("C-03", "G-01", "Match on R12",
     "Each payment line is looked up in Oracle by its 12-digit R12 number.",
     "The R12 makes PDF and Oracle numbers comparable.",
     "1. Make the R12 (G-01).\n2. Find it in the Open Item Report: brings bill-to, ship-to, date and open amount.\n"
     "3. Var (K) = paid - open. 0 = Matched.",
     "226860021923: paid 250.32, open 250.32, Matched. 186 of 242 lines found; 98 matched straight away.",
     "Automatic.", "-", "Regional sheets > Oracle Trx, Open Amount, Var (K)."),
    ("C-04", "Costco only", "Combine duplicate invoice numbers",
     "When the same invoice number appears twice, the lines are added together into one line.",
     "Costco shows a credit for the same invoice, day and store on a separate line. It belongs to that invoice; it is "
     "not a separate deduction.",
     "Group by R12 and add the amounts. The original amounts are kept in 'Combined Lines'.",
     "226860021923: 25,032.00 + (-24,781.68) = 250.32, matches Oracle 250.32.",
     "Automatic.", "-", "Formatted sheet; regional sheets > Combined Lines."),
    ("C-05", "G-03, G-04", "RV deductions",
     "References starting with RV are Costco deductions.",
     "The trade team works all Costco deductions through Ontario and allocates them later.",
     "Region Ontario. Uploaded under the RV number itself at amount x -1. Never written off, even under $11.",
     "RV000551013919 -8.75 -> uploaded +8.75 (debit) on Ontario 60477275.",
     "Automatic.", "Confirm small RVs are never written off (Q6).", "Ontario sheet (top); Upload."),
    ("C-06", "G-14, G-03", "Other deductions",
     "A number not found in Oracle that starts with 65, 14 or 604, or has 13 or more digits, is a deduction.",
     "These are trade programs and other deductions (Linda / Mehtab).",
     "Region Ontario, top of the section. Reference NF-<number> (13+ digits: NF-<R12>). Amount x -1.\n"
     "Important: the Oracle match is tried first. 65024 is in Oracle, so it is a normal match.",
     "6518087993 -43,591.80 -> NF-6518087993 +43,591.80. 1261732219883 -> NF-261732219883.",
     "Automatic.", "-", "Ontario sheet; Upload."),
    ("C-07", "Costco only", "Which region a line belongs to",
     "Each line is placed in Atlantic, Western, Ontario or Quebec.",
     "Each region has its own bill-to and its own NAC line.",
     "1. Bill-to of the Oracle document: 60455728 Atlantic, 60463822 Western, 60477275 Ontario, 60568255 Ontario "
     "(Additional), 60503143 Quebec.\n2. No Oracle document: by the number (leading zeros removed): 12 digits "
     "starting 1 = Atlantic, 2 = Western, 3 = Quebec; 11 or fewer / 13 or more = Ontario.\n"
     "3. Deductions always Ontario.",
     "026866321582 (11 digits without the 0) -> Ontario. 226603721414 -> Western.",
     "Automatic.", "If a Review note says the store belongs to another region, check it.", "Region column."),
    ("C-08", "Costco only", "Region letters", "Letters used in NAC names.", "Short region code in the reference.",
     "A = Atlantic, W = Western, O = Ontario, Q = Quebec.", "NAC1509518-W.", "Automatic.", "-", "Upload."),
    ("C-09", "G-14, G-03", "Invoice not in Oracle and nothing explains it: NF",
     "An invoice number not in Oracle, where the engine also finds no documents for that store and date.",
     "We must hold the money on the account until it is researched.",
     "NF-<number>, amount x -1.",
     "226603721414: 5,119.88 + (-10.88) = 5,109.00 paid -> NF-226603721414, upload -5,109.00.",
     "Automatic.", "-", "Regional sheet; Upload."),
    ("C-10", "Costco only", "Order within each region",
     "Lines are grouped so the sheet reads like the team's template.",
     "Easier to review; deductions on top of Ontario as today.",
     "Deductions > Matched > Resolved > WO > SP > OP > NF. Within a group, negatives first.",
     "", "Automatic.", "-", "Regional sheets."),
    ("C-11", "Costco only", "Look for a credit memo first",
     "If paid and open differ, look for an open credit memo of the same store, the same day and exactly the "
     "missing amount.",
     "Costco often reduces the payment by a credit memo that is already in Oracle.",
     "Found: attach it; the difference becomes 0 (Resolved).",
     "31 lines resolved this way in the sample.", "Automatic.", "-", "Resolution_Log (rule C-11)."),
    ("C-12", "G-02", "Try the neighbouring invoice numbers: x+1, x+2, x-1",
     "Costco often quotes a number next to the real invoice. If the paid invoice is x, the engine tries x+1, x+2 "
     "and x-1.",
     "In Oracle, x is often a credit memo and the real invoice is the number next to it.",
     "Accept a neighbour only if: same store (ship-to), same date, not used yet, and it makes the difference "
     "smaller. Exact 0 wins; otherwise the biggest improvement. Up to 2 neighbours per line.",
     "Store 656: paid 10,149.31 against 226307921597 (credit memo -557.03). x-1 = 226307921596 (10,806.62) is "
     "added. Open total 10,249.59, paid 10,149.31, L = +100.28 -> 656-080326-SP.",
     "Automatic. x-2 is not tried (add -2 to NEIGHBOUR_OFFSETS to enable).", "-",
     "Resolution_Log; 'Attached' rows on the regional sheets."),
    ("C-13", "G-02", "Combine 1 to 4 open items of the same store",
     "If a difference is left, the engine looks at all open items of that store on that day (then the day before, "
     "weekends included) and finds 1-4 of them that add up exactly to the difference.",
     "One Costco payment line can cover several invoices and credits.",
     "Smallest combination first. If more than one combination fits, the first is used and the line goes to Review.",
     "226860221501: paid 9,620.11, open 10,318.15, difference -698.04 = credits 226860221503 (-472.60) + "
     "226860221502 (-225.44). Resolved.",
     "Automatic.", "Check lines flagged 'more than one combination'.", "Resolution_Log (rule C-13)."),
    ("C-14", "Costco only", "Partial match",
     "If nothing explains the whole difference, one document of the same store and day can be attached when it "
     "leaves only a write-off-sized remainder (within $11).",
     "The team's practice is to reduce the variance as far as possible.",
     "Attached, remainder written off, line listed on Review.",
     "026866821574: remainder -3.75 written off.", "Automatic.", "Verify every partial match on Review.",
     "Review; Resolution_Log (rule C-14)."),
    ("C-15", "G-13", "Remainder of $11 or less is written off",
     "What is left after all matching (L) between -11.00 and +11.00 (inclusive) is a write-off.",
     "Costco's tolerance.", "Status WO; not uploaded; listed on WriteOffs.",
     "226603721412 L = 8.16; 026867021534 L = 6.72; 226860621945 L = 10.00. 6 write-offs in the sample.",
     "Automatic.", "Create the write-offs in Oracle after posting (G-13).", "WriteOffs."),
    ("C-16", "G-08, G-03, G-04", "Remainder above $11 is a short pay or overpay",
     "L above +11.00 = short pay (Costco paid less). L below -11.00 = overpay (Costco paid more).",
     "We keep a record of what Costco still owes, or of the extra money they sent.",
     "L = -(paid - open). Reference <store>-<MMDDYY of the invoice date>-SP / -OP; -1, -2 if already used. "
     "Upload amount = L (SP positive = debit, OP negative = credit). Above $1,000 the line also goes to Review.",
     "656-080326-SP +100.28. 542-081426-OP -37.92: paid 11,773.39 against open 11,735.47 (12,093.04 - 357.57).",
     "Automatic.", "Large SP/OP (e.g. 1768-080326-OP -4,538.78) usually mean a missing invoice or credit: research "
     "before upload.", "Regional sheets > Net Var (L), Reference; Review."),
    ("C-17", "G-05", "Regional upload with a NAC per region",
     "Each region's SP, OP, NF and deduction lines are balanced by NAC<payment>-<letter>.",
     "Each region is a different bill-to; each must balance to 0.",
     "NAC = -(sum of the region's upload lines). Write-offs excluded.",
     "NAC1509518-W = 14,143.75; -O = -120,183.06; -Q = 37.92 (no Atlantic lines this day).",
     "Automatic.", "-", "Upload (NAC Group column)."),
    ("C-18", "G-06, G-07, G-16", "Costco upload layout",
     "How Costco upload lines are filled.", "Matches the team's AutoInvoice practice.",
     "Sorted lowest to highest. Date = work date. Bill-to = the region's bill-to on every line.",
     "", "Automatic.", "Confirm payment terms for debit lines (Q18).", "Upload."),
    ("C-19", "G-09, G-15", "Costco receipt",
     "One receipt for the payment.", "Applies the cash.",
     "Lines: every matched or attached Oracle document at its open amount (12-digit numbers) + the regional NAC "
     "lines. Write-offs shown separately. Receipt + write-offs = remittance total.",
     "Receipt 1509518 = 898,467.62.", "Automatic.", "Confirm the receipt customer account (Q18).", "Receipt."),
    ("C-20", "G-08", "AutoInvoice rejects a duplicate reference (manual)",
     "Import AutoInvoice can reject a line whose reference already exists (e.g. a closed SP with the same name).",
     "The automation only sees open items.", "Fix in Manage AutoInvoice: add today's date to the reference.",
     "", "Not automated.", "Check the execution report after Import AutoInvoice.", "-"),
]

# ---------------------------------------------------------------------------------
# 6. Worked examples
# ---------------------------------------------------------------------------------
EXAMPLES = [
    ("Metro", "Normal invoice", "026593522226: gross 125.81, discount 0.31, net 125.50. Oracle 26593522226 open "
     "125.81.", "Numbers match after removing the 0; gross 125.81 = open 125.81.",
     "Matched. On the receipt as 026593522226 for 125.81.", "M-03, G-01, G-15"),
    ("Metro", "Credit already in Oracle", "026551222289: gross -112.00, no discount. Oracle open -112.00.",
     "-112.00 = -112.00.", "Matched. On the receipt for -112.00.", "M-03"),
    ("Metro", "Negative deduction", "0408274209: net -160,000.00, not in Oracle, 9 digits without the 0.",
     "Deduction (M-05). Not within +/-10. Upload = -160,000.00 x -1.",
     "Upload 408274209 +160,000.00, CACBC-DEDUC-MGMT-DR, Net 15.", "M-05, M-06, G-03, G-04"),
    ("Metro", "Positive deduction", "0408274235: net +16,148.06, not in Oracle.",
     "Deduction. Upload = 16,148.06 x -1.", "Upload 408274235 -16,148.06, CACBC-DEDUC-MGMT-CR.", "M-06"),
    ("Metro", "Payment of an unknown invoice", "26543919847: gross 1,430.03, net 1,426.45, not in Oracle, no "
     "closed-transactions export.", "Invoice-sized number, nothing known: NF (provisional). Upload = gross x -1 "
     "(DISC_BASIS TOTAL).", "NF-26543919847 -1,430.03 (credit); listed on Review to confirm in Oracle.",
     "M-09, G-14"),
    ("Metro", "Closed invoice, reference already used", "26921620976: gross 74.29. Oracle already has an open "
     "SB INV-26921620976 for 74.29.", "SB, and the name is taken, so -1.",
     "SB INV-26921620976-1, upload -74.29.", "M-10, G-08"),
    ("Metro", "Discount and NAC", "Deductions 41 lines: -572,219.34. NF/SB 7 lines gross: +2,373.45. Discount "
     "total 2,021.69.", "Upload: deductions +572,219.34, NF/SB -2,373.45, DISC +2,021.69 = 571,867.58. "
     "NAC = -571,867.58.", "DISC-0238539 +2,021.69; NAC0238539 -571,867.58; upload total 0.00.",
     "M-12, M-13, G-05"),
    ("Metro", "Receipt", "Matched invoices 736,853.60; NAC -571,867.58.", "736,853.60 - 571,867.58 = 164,986.02.",
     "Receipt NAC0238539 for 164,986.02 = remittance total.", "M-15, G-09"),
    ("Costco", "Invoice + Costco credit on two lines", "226860021923: 25,032.00 and -24,781.68 (same day, store "
     "56). Oracle open 250.32.", "Combine: 25,032.00 - 24,781.68 = 250.32 = open.", "Matched.", "C-04, C-03"),
    ("Costco", "Short pay via neighbour x-1 (store 656)", "226307921597: 10,249.59 and -100.28 -> paid 10,149.31. "
     "Oracle 226307921597 is a credit memo -557.03.", "Var = 10,149.31 - (-557.03) = 10,706.34. Neighbour x-1 "
     "226307921596 (10,806.62, same store/day) reduces it to -100.28. Nothing else fits. L = +100.28 > 11.",
     "SP 656-080326-SP +100.28 (debit, Western). Receipt applies 226307921597 (-557.03) and 226307921596 "
     "(10,806.62).", "C-12, C-16"),
    ("Costco", "Two credits explain the difference", "226860221501: paid 9,620.11, open 10,318.15, difference "
     "-698.04.", "Same store, same day: 226860221503 (-472.60) + 226860221502 (-225.44) = -698.04.",
     "Resolved (difference 0). All three documents on the receipt.", "C-13"),
    ("Costco", "Overpay (Quebec)", "326852522639: 11,814.19 and -40.80 -> paid 11,773.39. Oracle x is a credit "
     "-357.57; x-1 326852522638 = 12,093.04.", "Open 12,093.04 - 357.57 = 11,735.47. Paid 11,773.39 - 11,735.47 "
     "= 37.92 more. L = -37.92 < -11.", "OP 542-081426-OP -37.92 (credit). NAC1509518-Q = +37.92.",
     "C-12, C-16, C-17"),
    ("Costco", "Write-off", "226603721412: 1,650.84 and -8.16 -> paid 1,642.68. Oracle open 1,650.84.",
     "Nothing explains -8.16. L = +8.16, within +/-11.", "WO 8.16. Not uploaded. Write off in Oracle after posting.",
     "C-15, G-13"),
    ("Costco", "RV deduction", "RV000551013919: -8.75.", "RV = deduction, always Ontario, never written off. "
     "Upload = -8.75 x -1.", "Upload RV000551013919 +8.75 (DR) on 60477275.", "C-05"),
    ("Costco", "Deduction not in Oracle", "6518087993: -43,591.80. Starts with 65, not in Oracle.",
     "Deduction, Ontario, NF-<number>, amount x -1.", "NF-6518087993 +43,591.80 (DR).", "C-06"),
    ("Costco", "Invoice not in Oracle", "226603721414: 5,119.88 and -10.88 -> 5,109.00. Not in Oracle; nothing for "
     "that store and date.", "12 digits starting with 2 -> Western. NF, amount x -1.",
     "NF-226603721414 -5,109.00 (CR) in Western.", "C-07, C-09"),
    ("Costco", "Large overpay to research", "226713121547: 4,548.88 and -34.07 -> 4,514.81. Oracle x is a credit "
     "-23.97; no neighbour or combination fits.", "L = -(4,514.81 + 23.97) = -4,538.78.",
     "OP 1768-080326-OP -4,538.78, also on Review: the real invoice is probably missing from the open items.",
     "C-16"),
    ("Costco", "Engine result differs from the team (store 152)", "226305921535: paid 8,980.62, open 8,966.80, "
     "difference 13.82. The team booked OP -13.82.", "Same store and day: 226305921536 (+225.00) and 226305921537 "
     "(-211.18) add up to 13.82 exactly.", "Resolved by the engine. To be confirmed by the team (Q17).", "C-13"),
]

# ---------------------------------------------------------------------------------
# 7. Output guide
# ---------------------------------------------------------------------------------
OUTPUT = [
    ("Both", "Summary", "One-page overview with a description under every section: status banner, run information, "
     "reconciliation, how every line was classified, how the cash is explained (receipt and NAC), receipt summary and "
     "action items. Costco adds regions, automatic resolution and short pays / overpays.",
     "Read the banner first: green = ready, amber = ready after the Review sheet, red = a check failed."),
    ("Both", "Control_Checks", "Every automatic check with what it checks, PASS / FAIL, Expected / Actual / Difference, "
     "detail and what to do if it fails.",
     "Any FAIL: do not upload; find the cause."),
    ("Both", "Upload", "AutoInvoice lines. The first 14 columns follow the template; Category / Region / NAC Group / "
     "Source Invoice are helper columns.", "Totals 0 per NAC Group. Copy the template columns only, as text."),
    ("Both", "Receipt", "Receipt header on every row + one row per document applied (Reference Value / Amount).",
     "Sum of Reference Amount + write-offs = Receipt Amount."),
    ("Both", "WriteOffs", "Lines written off with their amount (Metro: WO Amount; Costco: WO Amount = paid - open).",
     "Create these write-offs in Oracle after posting the receipt."),
    ("Both", "Review", "Lines that need a human decision, with the reason in Note.",
     "Clear every line before uploading."),
    ("Metro", "Raw_Remittance", "The CSV lines exactly as read (audit copy).", "-"),
    ("Metro", "Formatted", "Every line with match result: Match Key, Oracle Trx / Amount / Bill-to, Variance, "
     "Status, Reference, Upload Amount, Receipt Amount, Rule, Note.", "Filter on Status to see each group."),
    ("Metro", "Deductions", "Deduction lines with net and upload amount.", "-"),
    ("Metro", "NF_SB", "NF and SB lines with the generated reference.", "Confirm provisional NFs in Oracle."),
    ("Costco", "Raw_Remittance", "All PDF lines with page number.", "-"),
    ("Costco", "Formatted", "Every payment line after combining duplicates, with its status (details per line on the "
     "regional sheets).", "-"),
    ("Costco", "Atlantic / Western / Ontario / Quebec", "Team-style regional sections: Invoice, PO, Date, Paid, "
     "Store, R12, Bill-to, Ship-to, Oracle Trx, Open Amount, Var (K), Net Var (L), Status, Reference, Upload Amount. "
     "Grey 'Attached' rows are extra documents the engine added to the line above.",
     "Status colours: green matched/resolved, yellow WO, orange SP, blue OP, red NF / Review, purple deduction."),
    ("Costco", "Resolution_Log", "One row per document attached by the engine: the rule used, the document, its "
     "amount and the variance before / after.", "Use it to understand any Resolved line."),
]

CHECKLIST = [
    ("1", "Before", "Run the Open Item Report (status Open) the same morning and save it in 'Open Item Reports'.",
     "G-11"),
    ("2", "Before", "Save the remittance (Metro CSV / Costco PDF) in the CODES folder.", "M-01, C-01"),
    ("3", "Before", "Close the rules workbook and old output files in Excel (the run cannot read / write open "
     "files).", "-"),
    ("4", "Run", "Metro: python automation/run_metro.py --remit \"<CSV>\" --open-items \"Open Item Reports/Open Item "
     "report Metro.xlsx\" (add --closed-trx <export> if you have one).", "-"),
    ("5", "Run", "Costco: python automation/run_costco.py --remit \"<PDF>\" --open-items \"Open Item Reports/open "
     "item report costco.xlsx\".", "-"),
    ("6", "Check", "Open the output in the Output folder. The Summary banner and the Control_Checks sheet must show all checks PASS.", "G-09, G-10"),
    ("7", "Check", "Work every line on the Review sheet (provisional NF, partial matches, engine matches on missing "
     "invoices, large SP/OP, duplicates).", "M-09, C-14, C-16"),
    ("8", "Check", "Spot-check a few Matched and Resolved lines against Oracle.", "C-11 to C-13"),
    ("9", "Oracle", "Copy the Upload template columns into the AutoInvoice template (as text) and run Load Interface "
     "File for Import > Import AutoInvoice. Fix any rejects (duplicate reference: add the date).", "G-12, C-20"),
    ("10", "Oracle", "Create the receipt from the Receipt sheet, post it, run AutoMatch.", "G-09"),
    ("11", "Oracle", "Query the receipt; if Amount Due is not 0, create the write-offs listed on WriteOffs "
     "(Bad Debt).", "G-13, M-17"),
    ("12", "After", "Save the output workbook with the remittance and share it with the team.", "-"),
]

PARAM_HELP = {
    "TRX_TYPE_DR": "Change only if Oracle's debit transaction type is renamed.",
    "TRX_TYPE_CR": "Change only if Oracle's credit transaction type is renamed.",
    "BUSINESS_UNIT": "Must match the business unit in the AutoInvoice template.",
    "BATCH_SOURCE": "Must match the transaction batch source in the AutoInvoice template.",
    "AMOUNT_EPSILON": "Leave as is (half a cent).",
    "SUFFIX_STYLE": "Use 'date' if -1/-2 suffixes keep clashing with closed transactions in Oracle.",
    "WO_TOLERANCE": "Change only if the company tolerance changes.",
    "WO_INCLUSIVE": "Y means the limit itself (e.g. exactly 10.00) is a write-off.",
    "GROUP_BILLTOS": "Add a bill-to here if Metro starts paying invoices of another related account.",
    "DISC_BASIS": "Switch to MATCHED if the team confirms DISC covers matched invoices only.",
    "RECEIPT_METHOD": "Set to the exact Oracle receipt method name.",
    "NEIGHBOUR_OFFSETS": "Add -2 to also try x-2.",
    "MAX_COMBINATION": "Higher values find bigger combinations but risk false matches; 4 is the team's maximum seen.",
    "REVIEW_RESIDUAL_ABOVE": "Lower it to review more SP/OP lines.",
    "NF_LEADING_ZERO": "Y to keep the leading 0 in NF references.",
    "TERMS_DR": "Payment terms for debit lines (Metro: Net 15).",
}

QUESTIONS = [
    ("Q1", "Metro: can we get a daily export of closed / all transactions?", "Without it, invoice numbers that are "
     "not open are marked NF provisionally and must be checked by hand to decide NF or SB.", "Open"),
    ("Q2", "Metro: is DISC the total discount (2,021.69) or only the discount on matched invoices (2,013.69)?",
     "Both give the same NAC; it changes the DISC amount and whether NF/SB use gross or net.", "Open (default total)"),
    ("Q3", "Metro: receipt accounts 60477295 (Ontario) / 60503130 (Quebec)?", "The walkthrough said 6047925 / "
     "60403130, which look like typos.", "Open"),
    ("Q6", "Costco: are small RV deductions always uploaded (never written off)?", "Oracle shows many RVs under $11 "
     "uploaded.", "Open (always uploaded)"),
    ("Q7", "Costco: not-found 12-digit numbers starting with 3 go to Quebec?", "The team rule only mentions 1 and 2.",
     "Open (Quebec)"),
    ("Q10", "Exact AutoInvoice template fixed values (business unit, batch source)?", "Business unit set to CACBC.",
     "Open"),
    ("Q13", "Metro: can one Ontario receipt be applied to Adonis / second Metro account invoices?",
     "31 of 1,380 matched invoices are on those accounts.", "Open"),
    ("Q14", "Exact Oracle receipt method name?", "Walkthrough mentions CB EFT / CC EFT JPM / CBFT.", "Open"),
    ("Q15", "Costco NF references with or without the leading 0?", "Team today: without (NF-26866321582).",
     "Open (without)"),
    ("Q16", "Costco: should x-2 also be tried?", "x+1, x+2 and x-1 are confirmed.", "Open"),
    ("Q17", "Costco store 152: engine match (two documents) or OP -13.82 as the team booked?", "See Worked_Examples.",
     "Open"),
    ("Q18", "Costco payment terms for debit lines and the receipt customer account?", "Currently blank / Ontario "
     "bill-to.", "Open"),
]


# ---------------------------------------------------------------------------------
def _write(wb, name, cols, rows, widths, common_col=None, section_col=None):
    ws = wb.create_sheet(name)
    ws.append(cols)
    for r in rows:
        ws.append(list(r))
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=Side(style="thin", color="808080"))
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = BORDER
        if common_col is not None:
            link = str(row[common_col].value or "")
            if link.startswith("G-") or link == "Both":
                for cell in row:
                    cell.fill = COMMON_FILL
                row[common_col].font = Font(bold=True)
        if section_col is not None:
            row[section_col].font = Font(bold=True)
    ws.freeze_panes = "B2" if len(cols) > 4 else "A2"
    return ws


def build(path=GUIDE_PATH):
    from config import load_params
    try:
        current = load_params()
    except SystemExit:
        current = {c: {n: d for cc, n, d, _ in DEFAULTS if cc == c} for c in {x[0] for x in DEFAULTS}}
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _write(wb, "Start_Here", ["Topic", "Explanation"], START, [30, 120], section_col=0)
    _write(wb, "Glossary", ["Term", "Meaning", "Example"], GLOSSARY, [32, 90, 40], section_col=0)
    _write(wb, "Process_Overview", ["Customer", "Step", "What happens", "Details", "Done by", "Where to see it",
                                    "Rules"], PROCESS, [10, 7, 30, 75, 16, 30, 18], section_col=2)
    _write(wb, "Decision_Flow", ["Customer", "Step", "Question", "If YES", "If NO", "Rules"], FLOW,
           [10, 7, 70, 45, 45, 14])
    rule_cols = ["Rule ID", "Common rule", "Rule", "What it means", "Why we do it", "How it works",
                 "Example from the sample files", "What the automation does", "What you need to check / do",
                 "Where to see it in the output"]
    widths = [8, 13, 30, 45, 45, 55, 55, 30, 40, 30]
    _write(wb, "Common_Rules", rule_cols, COMMON_RULES, widths, common_col=1, section_col=2)
    _write(wb, "Metro_Rules", rule_cols, METRO_RULES, widths, common_col=1, section_col=2)
    _write(wb, "Costco_Rules", rule_cols, COSTCO_RULES, widths, common_col=1, section_col=2)
    _write(wb, "Worked_Examples", ["Customer", "Scenario", "What the file / Oracle shows", "Calculation / logic",
                                   "Result", "Rules"], EXAMPLES, [10, 28, 55, 60, 50, 18], section_col=1)
    _write(wb, "Output_Guide", ["Workbook", "Sheet", "What it shows", "What to check"], OUTPUT, [10, 26, 90, 55],
           section_col=1)
    _write(wb, "Daily_Checklist", ["#", "Stage", "Action", "Rules"], CHECKLIST, [5, 10, 110, 16])
    prm = [(c, n, to_cell(current[c][n]), d, PARAM_HELP.get(n, "")) for c, n, _, d in DEFAULTS]
    ws = _write(wb, "Parameters_Explained", ["Customer", "Parameter", "Current value", "What it controls",
                                             "When would you change it?"], prm, [11, 24, 40, 70, 55])
    ws.insert_rows(1)
    ws["A1"] = ("Copy for reference. To change a value, edit the Parameters sheet in "
                "National_Customers_Automation_Rules.xlsx (the automation reads that file, not this one).")
    ws["A1"].font = Font(bold=True, italic=True)
    ws.freeze_panes = "A3"
    _write(wb, "Open_Questions", ["ID", "Question", "Why it matters", "Status"], QUESTIONS, [6, 70, 70, 20])
    wb.save(path)
    return path


if __name__ == "__main__":
    print(f"Rules guide written: {build()}")
