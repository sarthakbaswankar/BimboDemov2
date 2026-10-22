"""Builds 'Set of rules for automation/National_Customers_Automation_Rules.xlsx'.

Run once (or after editing the rules below):  python automation/build_rules_workbook.py
The Parameters sheet is read back by the automation at run time; edit values
there to change tolerances, bill-tos, transaction types, etc. Re-running this
script resets Parameters to the defaults in config.py.
"""
import sys
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

sys.path.insert(0, str(Path(__file__).resolve().parent))
from config import DEFAULTS, RULES_PATH, to_cell  # noqa: E402

RULE_COLS = ["Rule ID", "Step", "Rule", "Condition", "Action / Result", "Example (sample files)", "Source",
             "Automated?", "Notes"]

COMMON = [
    ("G-01", "Matching", "Invoice numbers are compared on a normalised key, never as typed",
     "Always", "Metro: digits without leading zeros. Costco: last 12 digits padded with 0 (R12)",
     "026866221906 = 26866221906; 1261732219883 -> 261732219883", "Costco help doc (R12); Metro transcript 45:09",
     "Y", "Replaces the R12 helper column and VLOOKUPs"),
    ("G-02", "Matching", "An open item can be applied only once per run", "Always",
     "Every document attached to a payment line is reserved", "", "Costco help doc ('don't use the same invoice twice')",
     "Y", ""),
    ("G-03", "Upload", "Upload amount = remittance amount x -1 (deductions, NF, SB); Costco SP/OP = net variance L",
     "Line goes to the upload", "Positive -> debit, negative -> credit",
     "Metro deduction 0152960350 net -55.06 -> +55.06", "Metro findings; handwritten notes p2", "Y", ""),
    ("G-04", "Upload", "Transaction type by sign", "Upload amount > 0 / < 0",
     "> 0: CACBC-DEDUC-MGMT-DR ; < 0: CACBC-DEDUC-MGMT-CR", "", "Costco help doc (corrected); handwritten notes p3",
     "Y", ""),
    ("G-05", "Upload", "Every upload balances to zero with a NAC line", "Per receipt (Metro) / per region (Costco)",
     "NAC amount = -(sum of the other upload lines)", "Metro NAC0238539 = -571,867.58",
     "Metro transcript 22:21 (clearing account = 0); Costco findings", "Y", ""),
    ("G-06", "Upload", "Date format YYYY-MM-DD", "All upload dates", "Metro: receipt (bank) date; Costco: work date",
     "2026-08-27", "Metro transcript 26:40", "Y", ""),
    ("G-07", "Upload", "Reference repeated in 3 fields; bill-to in 2 fields", "Every upload line",
     "Transaction number = line description = line flexfield segment 1; bill-to account = bill-to site", "",
     "Metro findings (Auto Invoice template)", "Y", "Fixed columns (BU, source, line type) come from Parameters"),
    ("G-08", "References", "References must be unique in Oracle", "Name already exists in Oracle or in this run",
     "Append -1, -2 ... (or -MMDDYY with SUFFIX_STYLE = date)", "SB INV-26921620976 exists -> SB INV-26921620976-1",
     "Metro transcript 12:02; handwritten notes p4; Costco transcript 51:34", "Y", ""),
    ("G-09", "Receipt", "Receipt references + write-offs = remittance total", "Always",
     "Control on the Controls sheet; FAIL blocks sign-off", "", "Metro findings; Costco help doc", "Y", ""),
    ("G-10", "Controls", "Run controls", "Every run",
     "Line total = header total; every line one status; upload = 0; receipt check; no double use; "
     "no duplicate reference; write-offs within tolerance", "", "Implementation plan 4.5", "Y", ""),
    ("G-11", "Inputs", "Open Item Report must be run the same morning, before processing", "Every run",
     "Report filtered to the right customers / collector, status Open", "", "Metro findings step 1; Costco findings",
     "N", "Manual prerequisite"),
    ("G-12", "Scope", "FBDI file generation and Oracle jobs are out of scope for now", "",
     "Output stops at the Upload and Receipt sheets", "", "Project brief", "N", "Later phase"),
    ("G-13", "Write-off", "Small differences are written off, not uploaded",
     "Amount within the customer tolerance (Metro +/-10, Costco +/-11, both inclusive)",
     "Not uploaded, not in NAC, listed on WriteOffs; Oracle: Create Write-off > Bad Debt after posting",
     "Costco 226603721412 L = 8.16", "Metro transcript 24:05; Costco help doc E; user confirmation", "Y",
     "Same rule, different tolerance per customer (see WriteOff_Rules)"),
    ("G-14", "NF", "Not-found items are booked as NF-<number>", "Number not in Oracle and not explained",
     "NF-<number>, upload amount x -1, never written off", "Metro NF-26506520282; Costco NF-6518087993",
     "Metro transcript 11:14; Costco help doc E", "Y", ""),
    ("G-15", "Receipt", "Invoice numbers on the receipt are 12 digits", "11-digit number from the Oracle export",
     "Leading 0 restored, kept as text", "26593522226 -> 026593522226", "Metro transcript 45:09; Costco help doc F",
     "Y", "Deduction references (RV, 65..., short Metro deductions) keep their own length"),
    ("G-16", "Upload", "Upload lines sorted by amount, lowest to highest", "Every upload",
     "Largest credit first, largest debit last", "", "Metro transcript 17:00; Costco help doc G", "Y", ""),
]

# Metro / Costco rules that apply a common rule (highlighted in their sheets)
METRO_COMMON = {
    "M-02": ["G-10"], "M-03": ["G-01", "G-02"], "M-06": ["G-03", "G-04"], "M-07": ["G-13"], "M-08": ["G-13"],
    "M-09": ["G-14", "G-03"], "M-10": ["G-08", "G-03"], "M-11": ["G-08"], "M-13": ["G-05"],
    "M-14": ["G-06", "G-07", "G-16"], "M-15": ["G-09"], "M-16": ["G-15"], "M-17": ["G-13"],
}
COSTCO_COMMON = {
    "C-02": ["G-10"], "C-03": ["G-01"], "C-05": ["G-03", "G-04"], "C-06": ["G-14", "G-03"],
    "C-09": ["G-14", "G-03"], "C-12": ["G-02"], "C-13": ["G-02"], "C-15": ["G-13"],
    "C-16": ["G-08", "G-03", "G-04"], "C-17": ["G-05"], "C-18": ["G-06", "G-07", "G-16"],
    "C-19": ["G-09", "G-15"], "C-20": ["G-08"],
}
# plain styling: light grey header, common rules lightly shaded
HEADER_FILL = PatternFill("solid", fgColor="F2F2F2")
HEADER_BORDER = Border(bottom=Side(style="thin", color="808080"))
COMMON_FILL = PatternFill("solid", fgColor="F7F7F7")

METRO = [
    ("M-01", "Input", "Remittance CSV with record types in column 1", "File received in the shared mailbox",
     "REMITTANCE SUMMARY = control/total; REMITTANCE = invoice lines; others captured separately",
     "CAN_RMT_METRODE-CA6-...: control 2000238539, total 164,986.02, 1,428 lines", "Metro transcript 2:58",
     "Y", "CA6 = Ontario, CA5 = Quebec; two separate payments (separate vendor numbers)"),
    ("M-02", "Input", "Remittance must tie", "Always", "Sum of net = summary total; gross - discount = net",
     "167,007.71 - 2,021.69 = 164,986.02", "Handwritten notes p1", "Y", "Moving Net before Gross is not needed"),
    ("M-03", "Matching", "Match on number without leading zeros against Metro-group bill-tos",
     "Number found as an open item", "Gross vs Oracle open amount; variance 0 -> Matched (applied on receipt)",
     "1,380 matched, 736,853.60, all variance 0", "Metro findings (columns I/J)", "Y",
     "Group = Metro Ontario, Metro Richelieu, Adonis (Parameters GROUP_BILLTOS)"),
    ("M-04", "Discount", "0.25% discount on invoices only; always use Net (column F), never Gross",
     "Invoice line", "Discount carried to the DISC memo", "Total discount 2,021.69", "Metro transcript 3:45, 21:07",
     "Y", ""),
    ("M-05", "Deductions", "N/A number with fewer than 10 significant digits = deduction",
     "Not in open items and < 10 digits (usually starts with 1 or 4)", "Status Deduction, reference = the number",
     "0152960344, 0408274209, 60122271 (41 deductions: 36 -ve, 5 +ve, -572,219.34)", "Metro transcript 8:21; notes p1",
     "Y", "Actual invoices are 10-12 digits"),
    ("M-06", "Deductions", "Deduction upload amount = net x -1", "Deduction outside +/-10",
     "+ve -> debit (Net 15), -ve -> credit", "0408274209 -160,000 -> +160,000 DR", "Metro transcript 14:14",
     "Y", "Deductions are uploaded whether or not they exist elsewhere in the system"),
    ("M-07", "Write-off", "Deduction within +/-$10 = write-off", "-10.00 <= net <= +10.00",
     "Not uploaded, not in NAC, listed on WriteOffs", "None in the 27-Aug file", "Metro transcript 9:20, 24:05",
     "Y", "Canada $10 tolerance"),
    ("M-08", "Write-off", "Matched invoice with a small variance", "|gross - open| <= 10",
     "Applied at Oracle amount; difference = write-off. Larger variance -> Review", "", "Plan 4.6", "Y", ""),
    ("M-09", "NF", "Invoice-shaped number not in Oracle = NF", "10+ digits, not open, not found in Oracle",
     "NF-<invoice>, upload x -1, never written off", "26506520282 1,693.27 -> NF-26506520282",
     "Metro transcript 11:14; notes p3", "Partial",
     "Without a closed-transactions export the NF is provisional and listed on Review"),
    ("M-10", "SB", "Invoice exists in Oracle but is closed = SB (should-be invoice)",
     "Found in closed-trx export, or SB INV-<n> already exists",
     "SB INV-<invoice> (+ -1/-2 if already used); +ve line -> credit, -ve line -> debit memo",
     "26921620976 74.29 -> SB INV-26921620976-1", "Metro transcript 11:41; notes p3-p4", "Partial", ""),
    ("M-11", "Duplicates", "Same invoice twice (remittance or already used in Oracle)", "Duplicate found",
     "Second occurrence handled as SB with -1/-2 suffix and listed on Review", "", "Metro transcript 12:02", "Y",
     "Usually a mistake, or a partial deduction/penalty"),
    ("M-12", "DISC", "Discount debit memo DISC-<last 7 of control>", "Discount > 0",
     "TOTAL basis: whole remittance discount, NF/SB at gross; MATCHED basis: matched discount, NF/SB at net",
     "DISC-0238539 = 2,021.69", "Handwritten notes p3", "Y", "Basis to be confirmed (Open_Questions)"),
    ("M-13", "NAC", "Balancing credit NAC<last 7 of control>", "Always", "NAC amount = -(deductions + NF/SB + DISC)",
     "NAC0238539 = -571,867.58", "Handwritten notes p5", "Y", ""),
    ("M-14", "Upload", "Upload layout", "", "Sorted largest negative to largest positive; positives Net 15; "
     "date = receipt date; Ontario and Quebec may be uploaded together", "", "Metro transcript 17:00, 27:23",
     "Y", ""),
    ("M-15", "Receipt", "Receipt contents", "",
     "Number = NAC name; amount = remittance total; account Ontario 60477295 / Quebec 60503130; "
     "batch 'Metro Ontario YYMMDD'; references = matched invoices (Oracle amount) + NAC; NF/SB excluded; "
     "sequence 1; type Transaction Number", "Receipt 164,986.02 = 736,853.60 - 571,867.58",
     "Metro transcript 34:03-37:55; notes p5", "Y", "Account numbers in the transcript (6047925 / 60403130) "
     "look like typos"),
    ("M-16", "Receipt", "Invoice numbers on the receipt are 12 digits", "11-digit number from the export",
     "Leading 0 restored (stored as text)", "26593522226 -> 026593522226", "Metro transcript 45:09", "Y",
     "Quebec invoices are 12 digits starting with 3"),
    ("M-17", "Oracle (manual)", "Post receipt, check Amount Due", "Amount Due <> 0",
     "Actions > Create Write-off > Bad Debt; then add the final (largest) invoice", "", "Metro transcript 40:54",
     "N", "Out of scope for now"),
    ("M-18", "Input", "Other record types (DSD ADJ, MARKET ADJUSTMENTS, COUPON)", "Data rows present",
     "Copied to Other_Records sheet for review", "None in the sample", "CSV layout", "Partial",
     "Treatment not described by the team yet"),
]

COSTCO = [
    ("C-01", "Input", "One payment advice PDF per day covering all regions", "Daily",
     "Weekday 6-10 pages, Monday 22-25 pages", "Payment 1509518, 9 pages, 251 lines", "Costco help doc A", "Y", ""),
    ("C-02", "Input", "PDF lines must equal Remittance TOTAL exactly", "Always", "Hard stop if not",
     "898,467.62", "Costco help doc B", "Y", "Replaces copy / text-to-columns / misplaced-column check"),
    ("C-03", "Matching", "Match on R12", "Number found as an open item",
     "Brings bill-to, ship-to, date, open amount; Var (K) = paid - open", "226860021923 250.32 = 250.32",
     "Costco help doc C, R12 notes", "Y", "186 of 242 lines found, 98 with variance 0"),
    ("C-04", "Combine", "Same invoice number twice = invoice + Costco's credit -> one line", "Duplicate number",
     "Amounts added, original amounts kept in 'Combined Lines'", "25,032.00 + (-24,781.68) = 250.32",
     "Costco findings 13:00", "Y", ""),
    ("C-05", "Deductions", "RV deductions", "Reference starts with RV",
     "Ontario, uploaded under the RV number at amount x -1, never written off", "RV000551013919 -8.75 -> +8.75",
     "Costco findings 16:25; Oracle open items (487 RVs on 60477275)", "Y", "Confirm no write-off for small RVs"),
    ("C-06", "Deductions", "Other deduction patterns", "Not found and starts with 65 / 14 / 604, or 13+ digits",
     "NF-<number> (13+ digits -> R12) in Ontario, at the top of the section", "6518087993 -> NF-6518087993 +43,591.80;"
     " 1261732219883 -> NF-261732219883", "Costco findings 16:00-17:35", "Y",
     "Matching runs first: 65024 is found in Oracle, so it is a match"),
    ("C-07", "Regions", "Region by bill-to; fallback by number pattern",
     "Bill-to known / unknown", "60455728 Atlantic, 60463822 Western, 60477275 + 60568255 Ontario, 60503143 Quebec."
     " Unknown: 12 digits starting 1 -> Atlantic, 2 -> Western, 3 -> Quebec; 11 or fewer / 13+ -> Ontario",
     "026866321582 -> Ontario", "Costco findings 15:00", "Y", "Pattern uses the number without leading zeros"),
    ("C-08", "Regions", "Region letters for NAC names", "", "A, W, O, Q", "NAC1509518-W", "Costco findings", "Y", ""),
    ("C-09", "NF", "Invoice not in Oracle and nothing found by the engine", "No document",
     "NF-<number>, amount x -1", "226603721414 5,109.00 -> NF-226603721414", "Costco help doc E", "Y",
     "Team convention drops the leading zero (NF_LEADING_ZERO)"),
    ("C-10", "Sorting", "Section order", "", "Deductions (Ontario), matched, resolved, WO, SP, OP, NF; "
     "negatives first within a group", "", "Costco findings 17:35", "Y", ""),
    ("C-11", "Resolution", "Credit memo first", "Variance <> 0",
     "Open credit memo with same ship-to, same date and amount = variance is attached", "",
     "Costco findings 22:46", "Y", ""),
    ("C-12", "Resolution", "Neighbour invoice x+1, x+2, x-1", "Variance <> 0",
     "Same ship-to and date, not used, must reduce |variance|; exact 0 preferred", "656: x-1 226307921596 -> SP 100.28",
     "Costco findings 24:30; user confirmation (x-1)", "Y", "x-2 not tried unless added to NEIGHBOUR_OFFSETS"),
    ("C-13", "Resolution", "Combination of up to 4 open items for the ship-to", "Variance still <> 0",
     "Same day first, then the day before (weekends included)", "226860221501: two docs = 698.04",
     "Costco findings 28:00-32:00", "Y", "Ambiguous combinations are flagged on Review"),
    ("C-14", "Resolution", "Partial match", "No exact set; one same-day document leaves a residual within +/-11",
     "Document attached, residual written off, flagged on Review", "", "Costco findings 29:00 ('reduce the variance')",
     "Y", ""),
    ("C-15", "Write-off", "Residual within +/-$11 = write-off", "-11.00 <= L <= +11.00",
     "Not uploaded, not in NAC, listed on WriteOffs", "226603721412 L = 8.16", "Costco help doc E; user confirmation",
     "Y", ""),
    ("C-16", "SP / OP", "Residual above +/-$11", "L = -(sum of K) ; L > 11 / L < -11",
     "SP <store>-<MMDDYY>-SP (debit) / OP <store>-<MMDDYY>-OP (credit); -1/-2 if reused",
     "656-080326-SP 100.28", "Costco findings 29:30", "Y", "Large residuals are also listed on Review"),
    ("C-17", "Upload", "Regional upload with balancing NAC", "Per region",
     "SP, OP, NF, deductions; WO excluded; NAC<payment>-<letter> = -(region total)", "", "Costco findings 37:20",
     "Y", ""),
    ("C-18", "Upload", "Upload layout", "", "Sorted lowest to highest; date = work date; bill-to = region bill-to",
     "", "Costco help doc G", "Y", ""),
    ("C-19", "Receipt", "Receipt contents", "", "Matched and engine-attached documents (12-digit invoice numbers) "
     "+ NAC lines; write-offs shown separately", "", "Costco help doc F", "Y", ""),
    ("C-20", "Oracle (manual)", "AutoInvoice reject on duplicate reference", "Import rejects a line",
     "Add today's date to the reference", "", "Costco transcript 51:34", "N", "Out of scope for now"),
]

WRITE_OFF = [
    ("Tolerance", "+/- $10", "+/- $11"),
    ("Boundary", "Inclusive: -10.00 <= amount <= +10.00", "Inclusive: -11.00 <= L <= +11.00"),
    ("Amount tested", "Net amount of a deduction (or variance of a matched invoice)",
     "Net variance L left after the resolution engine"),
    ("Above tolerance", "Deduction uploaded (x -1)", "L > 11 -> SP ; L < -11 -> OP"),
    ("Upload / NAC", "Not uploaded, not in NAC", "Not uploaded, not in NAC"),
    ("Receipt", "Shown separately; receipt refs + write-offs = remittance total", "Same"),
    ("Never written off", "NF / SB lines", "RV deductions, NF lines"),
    ("Oracle (manual)", "Create Write-off > Bad Debt after posting the receipt", "Same"),
    ("Example", "None in the 27-Aug file", "226603721412 8.16; 026867021534 6.72; 226860621945 10.00"),
]

NAMING = [
    ("NF-<invoice>", "Invoice / deduction not found in Oracle", "NF-26506520282 ; NF-6518087993", "Both"),
    ("SB INV-<invoice>", "Invoice exists in Oracle but is closed", "SB INV-26921620976-1", "Metro"),
    ("<deduction number>", "Metro deduction (< 10 digits)", "152960344", "Metro"),
    ("RV<number>", "Costco RV deduction", "RV000551013919", "Costco"),
    ("<store>-<MMDDYY>-SP", "Costco short pay", "656-080326-SP", "Costco"),
    ("<store>-<MMDDYY>-OP", "Costco overpay", "152-080326-OP", "Costco"),
    ("NAC<last 7 of control>", "Metro receipt number and balancing credit", "NAC0238539", "Metro"),
    ("DISC-<last 7 of control>", "Metro discount debit memo", "DISC-0238539", "Metro"),
    ("NAC<payment>-<A|W|O|Q>", "Costco regional balancing line", "NAC1509518-O", "Costco"),
    ("<ref>-1, -2", "Reference already exists", "656-080326-SP-1", "Both"),
]

QUESTIONS = [
    ("Q1", "Metro NF vs SB needs a closed/all-transactions export (--closed-trx)", "Open",
     "Without it NF is provisional and listed on Review"),
    ("Q2", "Metro DISC amount: 2,021.69 (total, NF/SB at gross) or 2,013.69 (matched, NF/SB at net)", "Open",
     "Default TOTAL (matches handwritten notes); Parameters DISC_BASIS"),
    ("Q3", "Metro receipt accounts 60477295 / 60503130 (transcript said 6047925 / 60403130)", "Open", "Using open items"),
    ("Q5", "Write-off boundaries", "Closed", "Metro +/-10 inclusive, Costco +/-11 inclusive"),
    ("Q6", "Costco small RV deductions always uploaded, never written off", "Open", "Implemented as always uploaded"),
    ("Q7", "Costco 12-digit not-found numbers starting with 3 -> Quebec", "Open", "Implemented"),
    ("Q10", "Exact FBDI / receipt template column names and fixed values (BU, source)", "Open",
     "Standard names used; Parameters BUSINESS_UNIT / BATCH_SOURCE"),
    ("Q11", "Treatment of DSD ADJ / MARKET ADJUSTMENTS / COUPON records", "Open", "Captured on Other_Records"),
    ("Q12", "Team's completed workbooks for the same days, for line-by-line comparison", "Open", ""),
    ("Q13", "Metro receipt applied to Adonis / second Metro bill-to invoices (31 lines)", "Open", "Flagged on Receipt"),
    ("Q14", "Exact Oracle receipt method name", "Open", "Parameters RECEIPT_METHOD"),
    ("Q15", "Costco NF with or without leading zero", "Open", "Default without (team convention)"),
    ("Q16", "Costco x-2 neighbour", "Open", "Not tried; add -2 to NEIGHBOUR_OFFSETS to enable"),
    ("Q17", "Costco store 152: engine resolves 226305921535 exactly with 226305921536 (+225.00) and "
     "226305921537 (-211.18); team booked OP -13.82", "Open", "Confirm which is correct"),
    ("Q18", "Costco payment terms for debit lines and receipt customer account", "Open",
     "Parameters TERMS_DR blank; receipt account = Ontario bill-to"),
]

# Costco rule changes found by cross validation (payment 1509518 vs the team's workbook, 22-Oct-2026).
# NOT in the code yet and not merged into Costco_Rules; see CODES/Costco_Rule_Changes.txt for the full detail.
PENDING_COLS = ["Change", "Rule affected", "New rule", "What the automation does now", "What the team does",
                "Example (payment 1509518)", "Where to change", "Status"]
COSTCO_PENDING = [
    ("P-01", "C-09 (and C-13 for not-found lines)",
     "Invoice not in Oracle (and not a deduction): book NF-<number> for the full paid amount straight away; "
     "do not search the store's other documents",
     "Searches the store's ship-tos and attaches other documents to the not-found invoice",
     "NF for the full amount, nothing attached",
     "226439924534 -> NF-226439924534 -9,779.25 Western (ours attached 2 documents). "
     "026866321582 -> NF-26866321582 -5,715.96 Ontario (ours attached 2 documents and wrote off -4.59)",
     "costco/engine.py run_costco: not-found branch sets status NF (no resolver call)",
     "Pending - not in code"),
    ("P-02", "C-13",
     "A combination may only be used if it does not overshoot: the documents can be attached in some order "
     "without the remaining gap ever getting bigger than the starting gap",
     "Accepts any 2-4 documents that add up to the gap, even if each is much bigger than the gap",
     "Books SP / OP when only oversized documents would close the gap",
     "Store 152, 226305921535: gap 13.82. Ours used +225.00 and -211.18; team booked 152-080326-OP -13.82. "
     "Answers Open_Questions Q17",
     "costco/resolver.py Resolver._subsets: add the no-overshoot check",
     "Pending - not in code"),
    ("P-03", "C-11 to C-14 (order)",
     "Exact first: credit memo, then one neighbour that closes the gap exactly, then an exact combination "
     "(a single document of the same ship-to/date counts). Only if nothing is exact: the neighbour that closes "
     "most of the gap, then search for an exact match again",
     "Takes the first neighbour that reduces the gap, before looking for an exact match",
     "Uses the document that closes the gap exactly, even if it is not x+1 / x+2 / x-1",
     "Store 1747, 226713021542: gap 3,546.12; x-4 = 3,546.12 exactly -> gap 0. Ours took x+1 and left "
     "1747-080326-OP -692.76",
     "costco/resolver.py Resolver.resolve: new order (exact search, then partial neighbours)",
     "Pending - not in code"),
    ("P-04", "C-12",
     "Also try neighbour x-2 (neighbours x+1, x+2, x-1, x-2)",
     "Tries x+1, x+2, x-1 only",
     "Uses x-2",
     "Store 1768, 226713121547: team attached x-2 226713121545 (3,581.42) -> 1768-080326-OP -957.36; "
     "ours left OP -4,538.78. Answers Open_Questions Q16",
     "Parameters NEIGHBOUR_OFFSETS = 1, 2, -1, -2 (also config.py default)",
     "Pending - not in code"),
    ("P-05", "C-06",
     "13+ digit deduction rule: count the digits after removing leading zeros",
     "Counts the digits as printed, so a leading 0 makes a 12-digit invoice look like a 13-digit deduction",
     "0265225211597 is the 12-digit invoice 265225211597 (starts with 2 -> Western)",
     "0265225211597 -> NF-265225211597 -4,257.61 Western (ours: deduction NF in Ontario)",
     "costco/engine.py is_deduction: compare the significant digits with DEDUCTION_MIN_DIGITS",
     "Pending - not in code"),
    ("Result", "All", "With P-01 to P-05 together, payment 1509518 matches the team's workbook completely "
     "(tested in a simulation, code not changed)",
     "Differs on 6 of 242 payment lines. NAC-O -120,183.06, NAC-W 14,143.75, write-offs 23.26",
     "NAC-O -118,724.71, NAC-W 23,920.25, NAC-Q 37.92, write-offs 27.85",
     "Simulated: Upload 66/66 and Receipt 284/284 lines identical, all control checks pass",
     "Apply in TESTING first, then in the main automation", "Pending"),
]

README = [
    ("Purpose", "Business rules applied by the Metro and Costco cash-application automation (automation/ folder)."),
    ("How it is used", "The Parameters sheet is read by the code on every run. Edit a value there (e.g. a tolerance "
     "or bill-to) to change behaviour without touching the code. Other sheets document the rules."),
    ("Run Metro", 'python automation/run_metro.py --remit "<Metro CSV>" --open-items "Open Item Reports/Open Item '
     'report Metro.xlsx" [--closed-trx <export>]'),
    ("Run Costco", 'python automation/run_costco.py --remit "<Costco PDF>" --open-items "Open Item Reports/open item '
     'report costco.xlsx" [--work-date YYYY-MM-DD]'),
    ("Output", "Output/Metro_<Region>_<Control>_<Date>.xlsx and Output/Costco_<Payment>_<Date>.xlsx"),
    ("Common rules", "Rows shaded light grey are common to Metro and Costco. In Metro_Rules and Costco_Rules the "
     "'Common Rule' column gives the matching G-xx rule in Common_Rules; Common_Rules lists where each rule is "
     "applied. Rows without shading are specific to that customer."),
    ("Costco pending changes", "Costco_Pending_Changes lists rule changes found by cross validation against the "
     "team's workbook (payment 1509518). They are NOT in the code yet and are not part of Costco_Rules until they "
     "are implemented. Full detail: CODES/Costco_Rule_Changes.txt"),
    ("Rebuild this file","python automation/build_rules_workbook.py (resets Parameters to code defaults)"),
    ("Version", f"Generated {date.today().isoformat()}"),
]


def _sheet(wb, name, cols, rows, widths):
    ws = wb.create_sheet(name)
    ws.append(cols)
    for r in rows:
        ws.append(list(r))
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill, cell.border = HEADER_FILL, HEADER_BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    return ws


def _highlight(ws, rows, bold_col=None):
    for r in rows:
        for cell in ws[r]:
            cell.fill = COMMON_FILL
        if bold_col:
            ws.cell(row=r, column=bold_col).font = Font(bold=True)


def build(path=RULES_PATH):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _sheet(wb, "Read_Me", ["Item", "Detail"], README, [20, 110])
    rule_widths = [8, 13, 40, 32, 48, 38, 30, 10, 36]

    # Common_Rules: where each common rule is applied
    def used_in(mapping, gid):
        return ", ".join(rid for rid, gids in mapping.items() if gid in gids) or "All (general rule)"

    common_rows = [r + (used_in(METRO_COMMON, r[0]), used_in(COSTCO_COMMON, r[0])) for r in COMMON]
    ws = _sheet(wb, "Common_Rules", RULE_COLS + ["Applied in Metro", "Applied in Costco"], common_rows,
                rule_widths + [18, 18])
    _highlight(ws, range(2, ws.max_row + 1))

    # Metro / Costco rules: 'Common Rule' column right after the Rule ID, common rows highlighted
    for name, rules, mapping in (("Metro_Rules", METRO, METRO_COMMON), ("Costco_Rules", COSTCO, COSTCO_COMMON)):
        rows = [(r[0], ", ".join(mapping.get(r[0], [])) or "-") + r[1:] for r in rules]
        ws = _sheet(wb, name, RULE_COLS[:1] + ["Common Rule"] + RULE_COLS[1:], rows,
                    rule_widths[:1] + [16] + rule_widths[1:])
        _highlight(ws, [i for i, r in enumerate(rules, start=2) if r[0] in mapping], bold_col=2)
    _sheet(wb, "Costco_Pending_Changes", PENDING_COLS, COSTCO_PENDING, [8, 18, 50, 40, 38, 55, 40, 16])
    _sheet(wb, "WriteOff_Rules", ["Topic", "Metro", "Costco"], WRITE_OFF, [20, 55, 55])
    _sheet(wb, "Parameters", ["Customer", "Parameter", "Value", "Description"],
           [(c, n, to_cell(v), d) for c, n, v, d in DEFAULTS], [11, 26, 45, 90])
    _sheet(wb, "Reference_Naming", ["Pattern", "Meaning", "Example", "Customer"], NAMING, [28, 45, 40, 10])
    _sheet(wb, "Open_Questions", ["ID", "Question", "Status", "Current handling"], QUESTIONS, [6, 80, 9, 55])
    wb.save(path)
    return path


if __name__ == "__main__":
    print(f"Rules workbook written: {build()}")
