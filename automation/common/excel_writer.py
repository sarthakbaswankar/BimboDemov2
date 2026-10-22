"""Writes the multi-sheet output workbook with consistent formatting."""
from datetime import date, datetime

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Colour scheme: navy headers, white rows with light row lines, and Excel's standard Good / Neutral / Bad
# colours on the Status cell only, so colour always means something:
#   green = done (Matched, Resolved, PASS)        amber = residual booked (WO, SP, OP)
#   red   = needs attention (NF, SB, Review, FAIL) blue  = posted as a new transaction (Deduction, DISC, NAC)
# Documents attached to a payment line are shown in grey italics under it.
NAVY = "1F3864"
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_BORDER = Border(bottom=Side(style="thin", color=NAVY))
SUB_HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
SUB_HEADER_FONT = Font(bold=True, color=NAVY)
ROW_BORDER = Border(bottom=Side(style="hair", color="BFBFBF"))
TOTAL_FILL = PatternFill("solid", fgColor="F2F2F2")
ATTACHED_FONT = Font(italic=True, color="808080")
LABEL_FONT = Font(bold=True, color="404040")
NOTE_FONT = Font(italic=True, color="595959")

GOOD, NEUTRAL, BAD, INFO = ("C6EFCE", "006100"), ("FFEB9C", "9C5700"), ("FFC7CE", "9C0006"), ("DDEBF7", NAVY)
STATUS_TAGS = {
    "Matched": GOOD, "Resolved": GOOD, "PASS": GOOD,
    "WO": NEUTRAL, "SP": NEUTRAL, "OP": NEUTRAL,
    "NF": BAD, "SB": BAD, "Review": BAD, "FAIL": BAD,
    "Deduction": INFO, "DISC": INFO, "NAC": INFO,
}
BANNER_TAGS = {"READY": GOOD, "REVIEW": NEUTRAL, "BLOCKED": BAD}
TAB_COLOURS = {"Summary": NAVY, "Control_Checks": "808080"}
MONEY_FORMAT = "#,##0.00;-#,##0.00"


def _value(v):
    if v is None:
        return None
    if isinstance(v, float) and v != v:
        return None
    if hasattr(v, "item") and not isinstance(v, (str, bytes)):  # numpy scalar
        try:
            return v.item()
        except (ValueError, AttributeError):
            pass
    return v


TOTAL_BORDER = Border(top=Side(style="thin", color=NAVY), bottom=Side(style="double", color=NAVY))
SUMMARY_SPAN = 9  # columns used by banners / section bars
# '#', Control, What it checks, Result, Expected, Actual, Difference, Detail, If it fails
CONTROL_SHEET_WIDTHS = [5, 40, 62, 9, 15, 15, 12, 40, 55]


def _tag(cell, status):
    """Colours one Status cell (fill + matching text), Excel Good/Neutral/Bad style."""
    tag = STATUS_TAGS.get(str(status or "").split(" ")[0])
    if tag:
        cell.fill = PatternFill("solid", fgColor=tag[0])
        cell.font = Font(bold=True, color=tag[1])
        cell.alignment = Alignment(horizontal="center", vertical=cell.alignment.vertical)


def _fmt_cell(cell, value, col_name=""):
    if isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"  # literal text, not a formula
    elif isinstance(value, float):
        cell.number_format = "0.0%" if str(col_name).startswith("%") else MONEY_FORMAT
    elif isinstance(value, (date, datetime)):
        cell.number_format = "yyyy-mm-dd"


def _write_blocks(ws, blocks, widths=None):
    """Sectioned summary sheet. Blocks:
    ('title', text) ('subtitle', text) ('banner', text, READY|REVIEW|BLOCKED) ('section', text)
    ('kv', [(label, value), ...]) ('table', df, {total: [cols], status_col: name}) ('note', text)"""
    span = max([SUMMARY_SPAN] + [len(b[1].columns) for b in blocks if b[0] == "table" and b[1] is not None])
    r = 1
    for b in blocks:
        kind = b[0]
        if kind == "title":
            ws.cell(row=r, column=1, value=b[1]).font = Font(bold=True, size=16, color=NAVY)
            r += 1
        elif kind == "subtitle":
            ws.cell(row=r, column=1, value=b[1]).font = NOTE_FONT
            r += 2
        elif kind == "banner":
            fill, text = BANNER_TAGS.get(b[2], INFO)
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=span)
            c = ws.cell(row=r, column=1, value=b[1])
            c.font = Font(bold=True, size=12, color=text)
            c.alignment = Alignment(wrap_text=True, vertical="center", indent=1)
            edge = Side(style="thin", color=text)
            for col in range(1, span + 1):
                cell = ws.cell(row=r, column=col)
                cell.fill = PatternFill("solid", fgColor=fill)
                cell.border = Border(top=edge, bottom=edge, left=Side(style="thick", color=text) if col == 1 else None,
                                     right=edge if col == span else None)
            ws.row_dimensions[r].height = 36
            r += 2
        elif kind == "section":
            for col in range(1, span + 1):
                ws.cell(row=r, column=col).fill = HEADER_FILL
            ws.cell(row=r, column=1, value=b[1]).font = Font(bold=True, size=12, color="FFFFFF")
            ws.row_dimensions[r].height = 20
            r += 1
        elif kind == "kv":
            for label, value in b[1]:
                ws.cell(row=r, column=1, value=label).font = LABEL_FONT
                c = ws.cell(row=r, column=2, value=_value(value))
                _fmt_cell(c, c.value)
                c.alignment = Alignment(horizontal="left")
                for col in (1, 2):
                    ws.cell(row=r, column=col).border = ROW_BORDER
                r += 1
            r += 1
        elif kind == "table":
            df, opts = b[1], (b[2] if len(b) > 2 else {})
            if df is None or not len(df):
                ws.cell(row=r, column=1, value="(none)").font = Font(italic=True, color="808080")
                r += 2
                continue
            cols = list(df.columns)
            for ci, name in enumerate(cols, start=1):
                c = ws.cell(row=r, column=ci, value=name)
                c.font, c.fill, c.border = SUB_HEADER_FONT, SUB_HEADER_FILL, HEADER_BORDER
                c.alignment = Alignment(wrap_text=True, vertical="center")
            r += 1
            s_idx = cols.index(opts["status_col"]) if opts.get("status_col") in cols else None
            for rec in df.itertuples(index=False):
                for ci, v in enumerate(rec, start=1):
                    c = ws.cell(row=r, column=ci, value=_value(v))
                    _fmt_cell(c, c.value, cols[ci - 1])
                    c.border = ROW_BORDER
                    if opts.get("wrap"):
                        c.alignment = Alignment(wrap_text=True, vertical="top")
                if s_idx is not None:
                    _tag(ws.cell(row=r, column=s_idx + 1), rec[s_idx])
                r += 1
            if opts.get("total"):
                ws.cell(row=r, column=1, value="Total").font = Font(bold=True, color=NAVY)
                for ci, name in enumerate(cols, start=1):
                    if name in opts["total"]:
                        tot = df[name].fillna(0).sum()
                        tot = round(float(tot), 2) if isinstance(tot, float) or str(df[name].dtype).startswith(
                            "float") else int(tot)
                        c = ws.cell(row=r, column=ci, value=tot)
                        _fmt_cell(c, tot, name)
                        c.font = Font(bold=True, color=NAVY)
                    ws.cell(row=r, column=ci).border = TOTAL_BORDER
                    ws.cell(row=r, column=ci).fill = TOTAL_FILL
                r += 1
            r += 1
        elif kind == "note":
            ws.cell(row=r, column=1, value=b[1]).font = NOTE_FONT
            r += 1
    widths = widths or [50, 62] + [16] * 12
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4


def write_workbook(path, sheets):
    """sheets: list of dicts {name, df, status_col (optional), money_cols (optional)}
    or {name, blocks} for a sectioned summary sheet."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for spec in sheets:
        if "blocks" in spec:
            ws = wb.create_sheet(spec["name"][:31])
            _write_blocks(ws, spec["blocks"], spec.get("widths"))
            if spec["name"] in TAB_COLOURS:
                ws.sheet_properties.tabColor = TAB_COLOURS[spec["name"]]
            continue
        df = spec["df"]
        ws = wb.create_sheet(spec["name"][:31])
        cols = list(df.columns)
        ws.append(cols)
        for c in range(1, len(cols) + 1):
            cell = ws.cell(row=1, column=c)
            cell.fill, cell.font, cell.border = HEADER_FILL, HEADER_FONT, HEADER_BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 30
        for rec in df.itertuples(index=False):
            ws.append([_value(v) for v in rec])
            for cell in ws[ws.max_row]:
                cell.border = ROW_BORDER
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    cell.data_type = "s"
        money = set(spec.get("money_cols") or [c for c in cols if str(df[c].dtype).startswith("float")])
        status_col = spec.get("status_col")
        s_idx = cols.index(status_col) + 1 if status_col in cols else None
        for ci, name in enumerate(cols, start=1):
            letter = get_column_letter(ci)
            width = len(str(name))
            for r in range(2, min(ws.max_row, 400) + 1):
                v = ws.cell(row=r, column=ci).value
                if v is not None:
                    width = max(width, len(f"{v:,.2f}") if isinstance(v, float) else len(str(v)))
            ws.column_dimensions[letter].width = min(max(width + 2, 9), 60)
            if name in money:
                for r in range(2, ws.max_row + 1):
                    ws.cell(row=r, column=ci).number_format = MONEY_FORMAT
            elif df[name].map(lambda x: isinstance(x, (date, datetime))).any() if len(df) else False:
                for r in range(2, ws.max_row + 1):
                    ws.cell(row=r, column=ci).number_format = "yyyy-mm-dd"
        if s_idx:
            for r in range(2, ws.max_row + 1):
                status = str(ws.cell(row=r, column=s_idx).value or "").split(" ")[0]
                if status == "Attached":
                    for c in range(1, len(cols) + 1):
                        ws.cell(row=r, column=c).font = ATTACHED_FONT
                else:
                    _tag(ws.cell(row=r, column=s_idx), status)
        ws.freeze_panes = "A2"
        if ws.max_row > 1 and cols:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{ws.max_row}"
    wb.save(path)
