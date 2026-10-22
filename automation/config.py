"""Run parameters.

DEFAULTS is the single source of truth for every tunable value. The rules
workbook builder writes these rows into the `Parameters` sheet, and at run
time `load_params()` reads that sheet back, so the AR team can change a value
in Excel without touching the code.
"""
from pathlib import Path

import openpyxl

CODES_DIR = Path(__file__).resolve().parent.parent
RULES_PATH = CODES_DIR / "Set of rules for automation" / "National_Customers_Automation_Rules.xlsx"

# (customer, parameter, default value, description)
DEFAULTS = [
    # ---- common -------------------------------------------------------------
    ("COMMON", "TRX_TYPE_DR", "CACBC-DEDUC-MGMT-DR", "Transaction type for positive upload amounts (debit)"),
    ("COMMON", "TRX_TYPE_CR", "CACBC-DEDUC-MGMT-CR", "Transaction type for negative upload amounts (credit)"),
    ("COMMON", "BUSINESS_UNIT", "CACBC","AutoInvoice business unit (fixed column; confirm exact name from team template)"),
    ("COMMON", "BATCH_SOURCE", "Import", "AutoInvoice transaction batch source (fixed column; confirm from team template)"),
    ("COMMON", "LINE_TYPE", "LINE", "AutoInvoice transaction line type"),
    ("COMMON", "CURRENCY", "CAD", "Currency code"),
    ("COMMON", "AMOUNT_EPSILON", 0.005, "Amounts closer than this are treated as equal (rounding)"),
    ("COMMON", "SUFFIX_STYLE", "number", "Duplicate reference suffix: 'number' (-1, -2) or 'date' (-MMDDYY)"),
    # ---- Metro --------------------------------------------------------------
    ("METRO", "WO_TOLERANCE", 10.00, "Write-off tolerance in $ (deduction net / matched variance)"),
    ("METRO", "WO_INCLUSIVE", True, "Y = exactly +/-10.00 is a write-off"),
    ("METRO", "INVOICE_MIN_DIGITS", 10, "N/A numbers with fewer significant digits than this are deductions"),
    ("METRO", "REGION_CA6", "Ontario", "Remittance file code CA6 -> region"),
    ("METRO", "REGION_CA5", "Quebec", "Remittance file code CA5 -> region"),
    ("METRO", "BILLTO_ONTARIO", "60477295", "Metro Ontario bill-to / receipt customer account"),
    ("METRO", "BILLTO_QUEBEC", "60503130", "Metro Quebec (Richelieu) bill-to / receipt customer account"),
    ("METRO", "GROUP_BILLTOS",
     ["60477295", "60463677", "60503130", "60453574", "60527663", "9201257", "60477237", "60503156"],
     "Bill-tos searched for matches (Metro Ontario, Metro Richelieu, Adonis). Others are never matched."),
    ("METRO", "TERMS_DR", "Net 15", "Payment terms for positive (debit) upload lines"),
    ("METRO", "TERMS_CR", "", "Payment terms for negative (credit) upload lines"),
    ("METRO", "NAC_PREFIX", "NAC", "Receipt / balancing credit memo prefix; NAC + last 7 digits of control"),
    ("METRO", "CONTROL_DIGITS", 7, "Number of trailing control-number digits used in NAC / DISC names"),
    ("METRO", "DISC_PREFIX", "DISC-", "Discount debit memo prefix"),
    ("METRO", "DISC_BASIS", "TOTAL",
     "TOTAL = DISC is the whole remittance discount and NF/SB upload at gross (handwritten notes). "
     "MATCHED = DISC is discount on matched invoices and NF/SB upload at net."),
    ("METRO", "NF_PREFIX", "NF-", "Not-found invoice prefix"),
    ("METRO", "SB_PREFIX", "SB INV-", "Should-be invoice prefix"),
    ("METRO", "RECEIPT_METHOD", "CB EFT JPM", "Receipt method (confirm exact Oracle name)"),
    ("METRO", "BATCH_NAME_PATTERN", "Metro {region} {yymmdd}", "Receipt batch name pattern"),
    # ---- Costco -------------------------------------------------------------
    ("COSTCO", "WO_TOLERANCE", 11.00, "Write-off tolerance in $ on residual net variance (L)"),
    ("COSTCO", "WO_INCLUSIVE", True, "Y = exactly +/-11.00 is a write-off"),
    ("COSTCO", "BILLTO_ATLANTIC", "60455728", "Atlantic bill-to (region letter A)"),
    ("COSTCO", "BILLTO_WESTERN", "60463822", "Western bill-to (region letter W)"),
    ("COSTCO", "BILLTO_ONTARIO", "60477275", "Ontario bill-to (region letter O); also receives all deductions"),
    ("COSTCO", "BILLTO_QUEBEC", "60503143", "Quebec bill-to (region letter Q)"),
    ("COSTCO", "BILLTO_ONTARIO_EXTRA", ["60568255"], "Additional bill-tos that belong to Ontario"),
    ("COSTCO", "DEDUCTION_PREFIXES", ["RV", "65", "14", "604"], "Not-found references starting with these are deductions"),
    ("COSTCO", "DEDUCTION_MIN_DIGITS", 13, "Not-found numbers with at least this many digits (as printed) are deductions"),
    ("COSTCO", "NEIGHBOUR_OFFSETS", [1, 2, -1], "Invoice-number neighbours tried (x+1, x+2, x-1)"),
    ("COSTCO", "MAX_COMBINATION", 4, "Max number of open items combined to explain one payment line"),
    ("COSTCO", "MAX_CANDIDATES", 30, "Max open items per ship-to/date considered in the combination search"),
    ("COSTCO", "SEARCH_DAY_BEFORE", True, "Also search the ship-to on the day before"),
    ("COSTCO", "REVIEW_RESIDUAL_ABOVE", 1000.00, "SP/OP residuals larger than this $ are also listed on the Review sheet"),
    ("COSTCO", "NF_LEADING_ZERO", False, "Y = NF-026866321582 ; N = NF-26866321582 (team's current convention)"),
    ("COSTCO", "TERMS_DR", "", "Payment terms for positive (debit) upload lines (confirm)"),
    ("COSTCO", "TERMS_CR", "", "Payment terms for negative (credit) upload lines"),
    ("COSTCO", "NAC_PATTERN", "NAC{payment}-{letter}", "Regional balancing line name"),
    ("COSTCO", "RECEIPT_METHOD", "CB EFT JPM", "Receipt method (confirm exact Oracle name)"),
]

REGION_LETTERS = {"Atlantic": "A", "Western": "W", "Ontario": "O", "Quebec": "Q"}


def _coerce(raw, default):
    """Convert a cell value from the Parameters sheet to the default's type."""
    if raw is None or raw == "":
        return "" if isinstance(default, str) else default
    if isinstance(default, bool):
        return str(raw).strip().upper() in ("Y", "YES", "TRUE", "1")
    if isinstance(default, list):
        items = [x.strip() for x in str(raw).split(",") if x.strip()]
        if default and isinstance(default[0], int):
            return [int(x) for x in items]
        return items
    if isinstance(default, int):
        return int(float(raw))
    if isinstance(default, float):
        return float(raw)
    if isinstance(raw, float) and raw.is_integer():
        raw = int(raw)
    return str(raw).strip()


def to_cell(value):
    """Default value -> text written into the Parameters sheet."""
    if isinstance(value, bool):
        return "Y" if value else "N"
    if isinstance(value, list):
        return ", ".join(str(v) for v in value)
    return value


def rules_label(rules_path=None):
    path = Path(rules_path) if rules_path else RULES_PATH
    return f"{path.name} (Parameters sheet)" if path.exists() else "built-in defaults (rules workbook not found)"


def load_params(rules_path=None):
    params = {}
    for cust, name, default, _ in DEFAULTS:
        params.setdefault(cust, {})[name] = default
    path = Path(rules_path) if rules_path else RULES_PATH
    if not path.exists():
        return params
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except PermissionError:
        raise SystemExit(f"Cannot read the rules workbook (it is probably open in Excel). "
                         f"Close it and run again:\n  {path}")
    if "Parameters" not in wb.sheetnames:
        return params
    defaults = {(c, n): d for c, n, d, _ in DEFAULTS}
    for row in wb["Parameters"].iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None or row[1] is None:
            continue
        cust, name = str(row[0]).strip().upper(), str(row[1]).strip()
        if (cust, name) in defaults:
            params[cust][name] = _coerce(row[2], defaults[(cust, name)])
    return params
