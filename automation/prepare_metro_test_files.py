"""Splits the team's completed Metro workbooks into automation inputs, by region.

For every workbook in the folder (e.g. 'METRO_ON 01-01-2026 $289422.50.xlsx'):
  * first sheet (the original remittance)  -> '<same name> CA6.csv' (Ontario) / '<same name> CA5.csv' (Quebec)
  * open items pasted on the 'Formatted' sheet (the block under the blue 'Trx Number' heading)
                                            -> '<same name> Open Item Report.xlsx'
Files go to <folder>/Segregated/Ontario (CA6) and <folder>/Segregated/Quebec (CA5).
The original workbooks are only read, never changed or moved. A Segregation_Log.xlsx lists every file,
its region, totals and any problem (e.g. a workbook open in Excel).

Run (from the CODES folder):
    python automation/prepare_metro_test_files.py "Metro/NEW FOR TEST"
"""
import argparse
import csv
import re
from collections import Counter
from datetime import date, datetime
from pathlib import Path

import openpyxl
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill

REGIONS = {"CA6": "Ontario (CA6)", "CA5": "Quebec (CA5)"}
OI_HEADER = ["Trx Number", "Trx Date", "Due Remaining", "Party Name", "Bill To Cust", "Ship To Loc", "Trx Number"]
HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")


def region_from_name(name):
    n = name.upper()
    if re.search(r"METRO_ON\b|\bONTARIO\b|_ON ", n):
        return "CA6"
    if re.search(r"QUEBEC|_QC\b|_QUE\b", n):
        return "CA5"
    return None


def csv_value(v):
    """Cell value -> CSV text in the same style as Metro's file (no scientific notation)."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y%m%d")
    if isinstance(v, float):
        if v.is_integer():
            return str(int(v))
        return repr(round(v, 6))
    return str(v).strip() if isinstance(v, str) else str(v)


def export_remittance(ws, out_path, max_cols=None):
    """First sheet -> CSV (max_cols limits the columns, e.g. 8 = A-H). Returns stats used for the log."""
    rows = []
    for r in ws.iter_rows(values_only=True):
        vals = [csv_value(v) for v in (r[:max_cols] if max_cols else r)]
        while vals and vals[-1] == "":
            vals.pop()
        rows.append(vals)
    while rows and not rows[-1]:
        rows.pop()
    with open(out_path, "w", newline="", encoding="utf-8") as fh:
        csv.writer(fh).writerows(rows)

    # stats: header row of REMITTANCE tells where Net is (the team moves Net before Gross)
    hdr = next((r for r in rows if r and r[0] == "REMITTANCE"), [])
    net_i = next((i for i, h in enumerate(hdr) if h.strip().lower() == "net"), None)
    lines = [r for r in rows if r and r[0] == "REMITTANCE" and r is not hdr]
    summary = [r for r in rows if r and r[0] == "REMITTANCE SUMMARY"][1:]
    net = sum(float(r[net_i]) for r in lines if net_i is not None and len(r) > net_i and r[net_i])
    totals = [float(r[4]) for r in summary if len(r) > 4 and r[4]]
    prefixes = Counter(r[4].zfill(12)[0] for r in lines if len(r) > 4 and r[4].isdigit() and len(r[4]) >= 11)
    return {"lines": len(lines), "controls": ", ".join(r[1] for r in summary), "net": round(net, 2),
            "total": round(sum(totals), 2), "column_order": " | ".join(hdr[5:8]), "prefixes": prefixes,
            "other_types": ", ".join(sorted({r[0] for r in rows[6:] if r and r[0] not in
                                             ("REMITTANCE", "REMITTANCE SUMMARY")}))}


def export_open_items(ws, out_path):
    """Open-item block under the 'Trx Number' heading on the Formatted sheet -> xlsx."""
    start_row = start_col = None
    data = []
    for i, r in enumerate(ws.iter_rows(values_only=True), start=1):
        if start_row is None:
            for c in range(len(r) - 2):
                if r[c] == "Trx Number" and r[c + 1] == "Trx Date" and r[c + 2] == "Due Remaining":
                    start_row, start_col = i, c
                    break
            continue
        block = list(r[start_col:start_col + 7]) + [None] * 7
        if block[0] not in (None, ""):
            data.append(block[:7])
    if start_row is None:
        return None
    wb = openpyxl.Workbook(write_only=True)
    out = wb.create_sheet("Sheet1")
    head = []
    for h in OI_HEADER:
        c = WriteOnlyCell(out, value=h)
        c.font, c.fill = Font(bold=True), HEADER_FILL
        head.append(c)
    out.append(head)
    for row in data:
        out.append(row)
    wb.save(out_path)
    return {"oi_rows": len(data), "oi_found_at": f"row {start_row}, column {openpyxl.utils.get_column_letter(start_col + 1)}"}


def process(folder):
    folder = Path(folder)
    out_root = folder / "Segregated"
    log = []
    for path in sorted(folder.glob("*.xlsx")):
        if path.name.startswith("~$"):
            continue
        entry = {"File": path.name, "Region": "", "Remittance CSV": "", "Open Item Report": "", "Status": "",
                 "Note": ""}
        code = region_from_name(path.stem)
        try:
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        except PermissionError:
            entry.update(Status="NOT DONE", Note="File is open in Excel; close it and run again")
            log.append(entry)
            continue
        first = wb.worksheets[0]
        notes = []
        if code is None:
            notes.append("Region not in file name")
        out_dir = out_root / REGIONS.get(code, "Unknown region")
        out_dir.mkdir(parents=True, exist_ok=True)
        csv_path = out_dir / f"{path.stem} {code or 'REGION'}.csv"
        stats = export_remittance(first, csv_path)
        source = first.title
        if stats["lines"] == 0 and "Formatted" in wb.sheetnames:
            # some workbooks have an empty first sheet; the original remittance is then in Formatted A-H
            stats = export_remittance(wb["Formatted"], csv_path, max_cols=8)
            source = "Formatted (columns A-H)"
            notes.append(f"First sheet '{first.title}' is empty; remittance taken from Formatted columns A-H")
        if stats["lines"] == 0:
            notes.append("No remittance lines found")
        # cross-check region with the invoice numbers (Ontario starts with 0, Quebec with 3)
        top = stats["prefixes"].most_common(1)
        content_code = {"0": "CA6", "3": "CA5"}.get(top[0][0]) if top else None
        if code and content_code and content_code != code:
            notes.append(f"Invoice numbers look like {content_code}, file name says {code}")
        if abs(stats["net"] - stats["total"]) > 0.01:
            notes.append(f"Net of lines {stats['net']:,.2f} does not equal remittance total {stats['total']:,.2f}")
        if stats["other_types"]:
            notes.append(f"Other record types present: {stats['other_types']}")
        oi = None
        if "Formatted" in wb.sheetnames:
            oi_path = out_dir / f"{path.stem} Open Item Report.xlsx"
            oi = export_open_items(wb["Formatted"], oi_path)
            if oi is None:
                notes.append("No 'Trx Number' heading found on Formatted sheet")
            else:
                entry["Open Item Report"] = oi_path.name
        else:
            notes.append("No Formatted sheet")
        wb.close()
        entry.update({
            "Region": REGIONS.get(code, "Unknown"), "Remittance CSV": csv_path.name,
            "Remittance taken from": source, "Control #": stats["controls"], "Remittance total": stats["total"],
            "Net of lines": stats["net"], "Remittance lines": stats["lines"],
            "Column order (F|G|H)": stats["column_order"],
            "Open item rows": oi["oi_rows"] if oi else None, "Open items found at": oi["oi_found_at"] if oi else "",
            "Status": "OK" if not notes else "CHECK", "Note": "; ".join(notes),
        })
        log.append(entry)
        print(f"{entry['Status']:8} {entry['Region']:14} {path.name}" + (f"  ({entry['Note']})" if notes else ""))
    for e in log:
        if e["Status"] == "NOT DONE":
            print(f"NOT DONE {'':14} {e['File']}  ({e['Note']})")
    write_log(out_root / "Segregation_Log.xlsx", log)
    return log


def write_log(path, log):
    cols = ["File", "Region", "Status", "Note", "Remittance taken from", "Control #", "Remittance total", "Net of lines",
            "Remittance lines", "Column order (F|G|H)", "Remittance CSV", "Open Item Report", "Open item rows",
            "Open items found at"]
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Log"
    ws.append(cols)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True), HEADER_FILL
    for e in log:
        ws.append([e.get(c) for c in cols])
    for i, w in enumerate([42, 15, 10, 50, 11, 13, 16, 16, 10, 22, 48, 58, 13, 18], start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    wb.save(path)


def main():
    ap = argparse.ArgumentParser(description="Split completed Metro workbooks into remittance CSV + open items")
    ap.add_argument("folder")
    process(ap.parse_args().folder)


if __name__ == "__main__":
    main()
