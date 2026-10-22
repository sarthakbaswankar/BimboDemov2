"""Renders the Summary sheet of output workbooks to HTML + PDF (no Excel needed).

The HTML is built from the actual cells of the generated workbook (values, number
formats, fills, fonts, merged cells, column widths), then printed to PDF with
Microsoft Edge in headless mode.

Example (from the CODES folder):
    python automation/preview_summary.py Output/Metro_Ontario_2000238539_20260827.xlsx Output/Costco_1509518_20260901.xlsx
"""
import argparse
import html
import subprocess
import tempfile
from datetime import date, datetime
from pathlib import Path

import openpyxl

EDGE_PATHS = [r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
              r"C:\Program Files\Microsoft\Edge\Application\msedge.exe"]
PX_PER_WIDTH_UNIT = 7
PAGE_WIDTH_PX = 960  # printable width of A4 landscape at 96 dpi, with a safety margin


def _rgb(color):
    if color is None or color.type != "rgb" or not isinstance(color.rgb, str):
        return None
    return "#" + color.rgb[-6:]


def _text(cell):
    v = cell.value
    if v is None:
        return "", None
    fmt = cell.number_format or ""
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        if "%" in fmt:
            return f"{v:.1%}", None
        if "#,##0.00" in fmt:
            return f"{v:,.2f}", ("#C00000" if v < 0 and "[Red]" in fmt else None)
        return f"{v:,}" if isinstance(v, int) and abs(v) >= 10000 else str(v), None
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d"), None
    return str(v), None


def sheet_to_html(path, sheet="Summary"):
    ws = openpyxl.load_workbook(path)[sheet]
    merged, covered = {}, set()
    for rng in ws.merged_cells.ranges:
        merged[(rng.min_row, rng.min_col)] = rng.max_col - rng.min_col + 1
        covered |= {(r, c) for r in range(rng.min_row, rng.max_row + 1)
                    for c in range(rng.min_col, rng.max_col + 1)} - {(rng.min_row, rng.min_col)}
    max_col = ws.max_column
    widths = [int((ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width or 9) * PX_PER_WIDTH_UNIT)
              for c in range(1, max_col + 1)]
    zoom = min(1.0, PAGE_WIDTH_PX / sum(widths))
    out = [f"<colgroup>{''.join(f'<col style=width:{w}px>' for w in widths)}</colgroup>"]
    for r in range(1, ws.max_row + 1):
        height = ws.row_dimensions[r].height
        cells = []
        for c in range(1, max_col + 1):
            if (r, c) in covered:
                continue
            cell = ws.cell(row=r, column=c)
            text, color = _text(cell)
            style = []
            if cell.fill is not None and cell.fill.fill_type == "solid" and _rgb(cell.fill.fgColor):
                style.append(f"background:{_rgb(cell.fill.fgColor)}")
            f = cell.font
            if f is not None:
                if f.bold:
                    style.append("font-weight:bold")
                if f.italic:
                    style.append("font-style:italic")
                if f.size and f.size != 11:
                    style.append(f"font-size:{f.size}pt")
                if not color and _rgb(f.color) and _rgb(f.color) != "#000000":
                    color = _rgb(f.color)
            if color:
                style.append(f"color:{color}")
            align = cell.alignment.horizontal if cell.alignment is not None else None
            if align in ("left", "center", "right"):
                style.append(f"text-align:{align}")
            elif isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                style.append("text-align:right")
            b = cell.border
            if b is not None:
                for side in ("top", "bottom", "left", "right"):
                    s = getattr(b, side)
                    if s is not None and s.style:
                        w = {"double": "3px double", "medium": "2px solid", "thick": "3px solid"}.get(s.style,
                                                                                                    "1px solid")
                        style.append(f"border-{side}:{w} {_rgb(s.color) or '#000'}")
            if cell.alignment is not None and cell.alignment.wrap_text:
                style.append("white-space:normal;vertical-align:top")
            span = f" colspan={merged[(r, c)]}" if (r, c) in merged else ""
            cells.append(f"<td{span} style=\"{';'.join(style)}\">{html.escape(text)}</td>")
        row_style = f" style=height:{int(height * 1.33)}px" if height else ""
        out.append(f"<tr{row_style}>{''.join(cells)}</tr>")
    title = html.escape(str(ws.cell(row=1, column=1).value or sheet))
    return f"""<!doctype html><html><head><meta charset="utf-8"><title>{title}</title><style>
@page {{ size: A4 landscape; margin: 10mm; }}
body {{ font-family: Calibri, 'Segoe UI', Arial, sans-serif; font-size: 11pt; margin: 0; zoom: {zoom:.3f}; }}
table {{ border-collapse: collapse; table-layout: fixed; }}
td {{ padding: 1px 4px; white-space: nowrap; overflow: visible; vertical-align: middle; }}
</style></head><body><table style="width:{sum(widths)}px">{''.join(out)}</table></body></html>"""


def to_pdf(html_path, pdf_path):
    """Returns None if OK, otherwise the reason the PDF was not created."""
    edge = next((p for p in EDGE_PATHS if Path(p).exists()), None)
    if not edge:
        return "Microsoft Edge not found"
    for _attempt in range(2):  # Edge occasionally hangs; stop it and try once more
        # separate throw-away profile so a running Edge window does not block the headless print
        with tempfile.TemporaryDirectory(prefix="edgepdf_", ignore_cleanup_errors=True) as profile:
            # no pipes: Edge's helper processes would keep them open and block
            proc = subprocess.Popen([edge, "--headless=new", "--disable-gpu", "--no-pdf-header-footer",
                                     f"--user-data-dir={profile}", f"--print-to-pdf={pdf_path}",
                                     Path(html_path).resolve().as_uri()],
                                    stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
            try:
                proc.wait(timeout=60)
            except subprocess.TimeoutExpired:
                proc.kill()
        if Path(pdf_path).exists():
            return None
    return "Edge did not produce the PDF (try again)"


def main():
    ap = argparse.ArgumentParser(description="Render output Summary sheets to HTML/PDF")
    ap.add_argument("workbooks", nargs="+")
    ap.add_argument("--sheet", default="Summary")
    a = ap.parse_args()
    for wb in a.workbooks:
        wb = Path(wb).resolve()  # Edge needs absolute paths
        name = f"{a.sheet}_Sample_" + wb.stem.split("_")[0]
        html_path, pdf_path = wb.with_name(name + ".html"), wb.with_name(name + ".pdf")
        html_path.write_text(sheet_to_html(wb, a.sheet), encoding="utf-8")
        if pdf_path.exists():
            pdf_path.unlink()
        err = to_pdf(html_path, pdf_path)
        print(f"{wb.name}: {html_path.name}" + (f", {pdf_path.name}" if not err else f" (PDF not created: {err})"))


if __name__ == "__main__":
    main()
