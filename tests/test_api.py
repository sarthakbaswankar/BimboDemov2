"""API tests: python -m pytest tests -q   (from the 'demo National Customers' folder)"""
import base64
import io
import sys
import tempfile
import zipfile
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import api.index as api_module  # noqa: E402
from api.index import app  # noqa: E402

api_module.LOCAL_OUTPUT = Path(tempfile.mkdtemp(prefix="api_test_output_"))   # keep the real Output folder clean
client = TestClient(app)
METRO = ROOT / "CAN_RMT_METRODE-CA6-RMT_2252_202608271334.csv"
COSTCO = ROOT / "Costco Payment.PDF"


def post(customer, path, filename=None, uploaded_as=None, **extra):
    with open(path, "rb") as fh:
        return client.post("/process_national_customers/upload", data={"customer": customer, "filename": filename or path.name, **extra},
                           files={"file": (uploaded_as or path.name, fh)})


def test_info_and_health():
    assert client.get("/health").json() == {"status": "ok"}
    info = client.get("/").json()
    assert info["customers"] == ["costco", "metro"]
    assert all(v["present"] for v in info["open_item_reports"].values())


def test_metro_excel():
    r = post("metro", METRO)
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert r.headers["x-controls"] == "ALL PASS"
    assert "Metro_Ontario_2000238539_20260827.xlsx" in r.headers["content-disposition"]
    wb = openpyxl.load_workbook(io.BytesIO(r.content), read_only=True)
    assert {"Summary", "Control_Checks", "Upload", "Receipt"} <= set(wb.sheetnames)
    assert wb.sheetnames[0] == "Summary" and wb.sheetnames[-1] == "Control_Checks"


def test_metro_json():
    r = post("metro", METRO, output="json")
    res = r.json()["results"][0]
    assert res["remittance_total"] == 164986.02 and res["nac_amount"] == -571867.58 and res["controls"] == "ALL PASS"


def test_costco_excel():
    r = post("costco", COSTCO)
    assert r.status_code == 200, r.text
    assert r.headers["x-controls"] == "ALL PASS"
    assert "Costco_1509518_20260901.xlsx" in r.headers["content-disposition"]
    wb = openpyxl.load_workbook(io.BytesIO(r.content), read_only=True)
    assert wb.sheetnames[0] == "Summary" and wb.sheetnames[-1] == "Control_Checks"


@pytest.mark.parametrize("customer,path,expect", [("Costco", COSTCO, "Costco_1509518_20260901.xlsx"),
                                                  ("Metro", METRO, "Metro_Ontario_2000238539_20260827.xlsx")])
def test_json_payload(customer, path, expect):
    body = {"customer": customer, "filename": path.name,
            "file_base64": base64.b64encode(path.read_bytes()).decode()}
    default = client.post("/process_national_customers", json=body)            # no 'output' -> JSON with base64
    assert default.headers["content-type"].startswith("application/json")
    assert default.json()["output_file_name"] == expect and default.json()["file_base64"]
    r = client.post("/process_national_customers", json={**body, "output": "excel"})
    assert r.status_code == 200, r.text
    assert r.headers["x-controls"] == "ALL PASS" and expect in r.headers["content-disposition"]
    summary = client.post("/process_national_customers", json={**body, "output": "json"}).json()
    assert summary["results"][0]["output_file_name"] == expect
    assert r.headers["x-saved-to"] == f"Output/{expect}" and (api_module.LOCAL_OUTPUT / expect).exists()


def test_json_payload_bad_base64():
    r = client.post("/process_national_customers", json={"customer": "costco", "filename": "x.pdf", "file_base64": "not base64!"})
    assert r.status_code == 400
    assert r.json() == {"status": "error", "code": 400, "detail": "file_base64 is not valid base64."}


def test_base64_response_for_visual_builder():
    body = {"customer": "Costco", "filename": COSTCO.name, "file_base64": base64.b64encode(COSTCO.read_bytes()).decode(),
            "output": "base64"}
    res = client.post("/process_national_customers", json=body).json()
    assert res["status"] == "success" and res["controls"] == "ALL PASS"
    assert res["output_file_name"] == "Costco_1509518_20260901.xlsx" and res["content_type"].endswith(".sheet")
    wb = openpyxl.load_workbook(io.BytesIO(base64.b64decode(res["file_base64"])), read_only=True)
    assert wb.sheetnames[0] == "Summary"


def test_missing_field_error_shape():
    r = client.post("/process_national_customers", json={"customer": "costco"})
    assert r.status_code == 422 and r.json()["status"] == "error" and "file_base64" in r.json()["detail"]


def test_cors_preflight_from_browser_app():
    r = client.options("/process_national_customers", headers={
        "Origin": "https://example.oraclecloud.com", "Access-Control-Request-Method": "POST",
        "Access-Control-Request-Headers": "content-type,x-api-key"})
    assert r.status_code == 200 and r.headers["access-control-allow-origin"] in ("*", "https://example.oraclecloud.com")


@pytest.mark.parametrize("customer,path", [("METRO", METRO), ("Metro", METRO), (" Costco ", COSTCO),
                                           ("COSTCO", COSTCO)])
def test_customer_name_any_case(customer, path):
    r = post(customer, path, output="json")
    assert r.status_code == 200, r.text
    assert r.json()["customer"] == customer.strip().lower()
    assert r.json()["results"][0]["controls"] == "ALL PASS"


def test_open_items_kept_in_memory_and_runs_repeat():
    first = post("metro", METRO, output="json").json()["results"][0]
    again = post("metro", METRO, output="json").json()["results"][0]   # matched against the in-memory report
    assert first == again and again["nac_amount"] == -571867.58
    reports = client.get("/").json()["open_item_reports"]
    assert reports["metro"]["rows_in_memory"] and reports["costco"]["rows_in_memory"]


@pytest.mark.parametrize("customer,path,typed,region", [
    ("metro", METRO, "CAN_RMT_METRODE-CA6-RMT_2252_202608271334", "Ontario"),   # typed name without .csv
    ("metro", METRO, "Metro remittance.csv", "Ontario"),    # CA6 taken from the uploaded file's own name
    ("metro", METRO, "metro", "Ontario"),
    ("costco", COSTCO, "Costco Payment", None),             # Costco: the PDF is enough, any name
    ("costco", COSTCO, "costco", None),
])
def test_names_handled_like_main_automation(customer, path, typed, region):
    r = post(customer, path, filename=typed, output="json")
    assert r.status_code == 200, r.text
    res = r.json()["results"][0]
    assert res["controls"] == "ALL PASS"
    if region:
        assert res["output_file_name"].startswith(f"Metro_{region}_") and res["nac_amount"] == -571867.58


@pytest.mark.parametrize("customer,path,typed,uploaded_as,status", [
    ("walmart", METRO, "x.csv", None, 400),                       # unknown customer
    ("costco", METRO, "x.pdf", None, 400),                        # Costco but the file is a CSV
    ("metro", COSTCO, "CA6.csv", None, 400),                      # Metro but the file is a PDF
    ("metro", METRO, "remittance.csv", "remittance.csv", 422),    # no CA5/CA6 in either name
])
def test_bad_requests(customer, path, typed, uploaded_as, status):
    r = post(customer, path, filename=typed, uploaded_as=uploaded_as)
    assert r.status_code == status, r.text


def test_swagger_placeholder_values_ignored():
    # /docs sends the text 'string' in optional fields left empty; Costco must still work
    r = post("costco", COSTCO, filename="Costco Payment.PDF", region="string", output="string")
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")


@pytest.mark.parametrize("url", ["/api/index?__vpath=/health",   # Vercel hands over the rewrite destination
                                 "/health?__vpath=/health",      # Vercel hands over the original path
                                 "/health"])                     # local run, no rewrite
def test_vercel_rewrite_paths(url):
    assert client.get(url).json() == {"status": "ok"}


def test_vercel_rewrite_root_docs_and_post():
    assert client.get("/api/index?__vpath=/").json()["service"].startswith("Bimbo")
    assert client.get("/api/index?__vpath=/openapi.json").status_code == 200
    assert client.get("/api/index?__vpath=/docs").status_code == 200
    body = {"customer": "Costco", "filename": COSTCO.name, "file_base64": base64.b64encode(COSTCO.read_bytes()).decode()}
    res = client.post("/api/index?__vpath=/process_national_customers", json=body).json()
    assert res["status"] == "success" and res["controls"] == "ALL PASS"
